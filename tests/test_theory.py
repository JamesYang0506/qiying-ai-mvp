import hashlib
import io
import unittest
import zipfile

from openpyxl import Workbook, load_workbook
from profit_model.theory import normalize_baseline, preview_theory, theory_template


def workbook(rows, percent_cell=None):
    book = Workbook()
    sheet = book.active
    sheet.title = '理论利润表'
    for row in rows:
        sheet.append(row)
    if percent_cell:
        sheet[percent_cell].number_format = '0%'
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


class TheoryTests(unittest.TestCase):
    def quote(self, **updates):
        data = dict(
            method='quote', business_order_id='业务001', quantity='100500',
            unit_price='17.2', material_pct='65', production_management_pct='2.8',
            production_loss_pct='0.8', defect_loss_pct='1.8', energy_pct='0.2',
            preparation_labor_per_pair='0', press_labor_per_pair='0',
            company_management_pct='2', administration_pct='1', tax_basis='ignored',
        )
        data.update(updates)
        return data

    def test_quote_zero_labor_and_expenses_classification(self):
        result = normalize_baseline(self.quote())
        self.assertEqual(result['revenue'], '1728600')
        self.assertEqual(result['production_cost'], '1220391.6')
        self.assertEqual(result['target_margin_pct'], '29.4')
        self.assertEqual(result['operating_expense'], '51858')
        self.assertEqual(result['operating_profit'], '456350.4')
        self.assertEqual(result['labor_cost'], '48400.8')
        self.assertEqual(result['manufacturing_cost'], '3457.2')
        self.assertEqual(result['other_production_cost'], '44943.6')
        self.assertEqual(result['tax_basis'], 'ignored')
        self.assertIsNone(result['target_defect_pct'])
        self.assertEqual(result['defect_loss_pct'], '1.8')

    def test_quote_labor_required_positive_price_and_individual_percentage_limits(self):
        for field in ('preparation_labor_per_pair', 'press_labor_per_pair'):
            for missing in (None, ''):
                with self.subTest(field=field, missing=missing), self.assertRaises(ValueError):
                    normalize_baseline(self.quote(**{field: missing}))
        for updates in (
            {'quantity': '0'}, {'unit_price': '0'}, {'energy_pct': '101'},
            {'material_pct': '-1'}, {'preparation_labor_per_pair': '-1'},
        ):
            with self.subTest(updates=updates), self.assertRaises(ValueError):
                normalize_baseline(self.quote(**updates))
        result = normalize_baseline(self.quote(material_pct='100'))
        self.assertLess(float(result['target_margin_pct']), 0)

    def test_quote_labor_and_management_do_not_change_production_scope(self):
        result = normalize_baseline(
            self.quote(preparation_labor_per_pair='1', press_labor_per_pair='.5')
        )
        self.assertEqual(result['target_margin_pct'], '20.679069767442')
        self.assertEqual(result['labor_cost'], '199150.8')
        result = normalize_baseline(
            self.quote(company_management_pct='0', administration_pct='0')
        )
        self.assertEqual(result['target_margin_pct'], '29.4')

    def test_quote_excel_preview_accepts_standard_fields_and_preserves_sources(self):
        rows = [['字段', '数值']] + [[key, value] for key, value in self.quote().items()]
        result = preview_theory(workbook(rows))
        self.assertEqual(result['baseline']['target_margin_pct'], '29.4')
        self.assertEqual(result['baseline']['production_management_pct'], '2.8')
        self.assertTrue(any('物流' in warning for warning in result['warnings']))

    def test_quote_all_source_fields_required_and_percent_formats(self):
        data = self.quote()
        for field in set(data) - {'method', 'business_order_id', 'tax_basis'}:
            missing = dict(data)
            del missing[field]
            with self.subTest(field=field), self.assertRaisesRegex(ValueError, '不能为空'):
                normalize_baseline(missing)
        data['material_pct'] = '65%'
        data['defect_loss_pct'] = '1.8%'
        result = normalize_baseline(data)
        self.assertEqual(result['target_margin_pct'], '29.4')
        rows = [['字段', '数值']] + [[key, value] for key, value in self.quote().items()]
        material_row = next(i for i, row in enumerate(rows, 1) if row[0] == 'material_pct')
        rows[material_row - 1] = ['材料成本比例', .65]
        result = preview_theory(workbook(rows, f'B{material_row}'))
        self.assertEqual(result['baseline']['material_pct'], '65')

    def test_expanded_template_remains_valid_for_ratio_with_quote_fields_empty(self):
        book = load_workbook(io.BytesIO(theory_template()))
        sheet = book['理论利润表']
        for row in sheet.iter_rows(min_row=2):
            if row[0].value == '业务订单号':
                row[1].value = '业务001'
            elif row[0].value == '目标生产毛利率':
                row[1].value = 20
        stream = io.BytesIO()
        book.save(stream)
        result = preview_theory(stream.getvalue())
        self.assertEqual(result['baseline']['target_margin_pct'], '20')

    def ratio(self, **updates):
        return dict(
            method='ratio',
            business_order_id='业务001',
            target_margin_pct='20',
            **updates,
        )

    def costs(self, **updates):
        data = dict(
            method='costs',
            business_order_id='业务001',
            quantity='100',
            unit_price='10',
            discount_per_pair='1',
            material_cost='400',
            labor_cost='100',
            manufacturing_cost='100',
            other_production_cost='120',
        )
        data.update(updates)
        return data

    def test_ratio_preserves_missing_costs_and_negative_target(self):
        result = normalize_baseline(self.ratio())
        self.assertEqual(result['target_margin_pct'], '20')
        self.assertEqual(result['unit'], '双')
        self.assertIsNone(result['revenue'])
        self.assertIsNone(result['production_cost'])
        self.assertIsNone(result['material_cost'])
        self.assertEqual(
            normalize_baseline(dict(self.ratio(), target_margin_pct='-5'))[
                'target_margin_pct'
            ],
            '-5',
        )

    def test_costs_calculated_independently(self):
        result = normalize_baseline(self.costs())
        self.assertEqual(
            [
                result[k]
                for k in (
                    'revenue',
                    'production_cost',
                    'production_profit',
                    'target_margin_pct',
                )
            ],
            ['900', '720', '180', '20'],
        )
        self.assertEqual(
            normalize_baseline(self.costs(target_margin_pct='20.01'))[
                'target_margin_pct'
            ],
            '20',
        )
        with self.assertRaisesRegex(ValueError, '冲突'):
            normalize_baseline(self.costs(target_margin_pct='20.02'))

    def test_required_fields_and_numeric_limits(self):
        for field in (
            'quantity',
            'unit_price',
            'discount_per_pair',
            'material_cost',
            'labor_cost',
            'manufacturing_cost',
            'other_production_cost',
        ):
            data = self.costs()
            del data[field]
            with self.subTest(field=field), self.assertRaises(ValueError):
                normalize_baseline(data)
        for value in (
            'NaN',
            'Infinity',
            '1e999999999',
            '1000000000000001',
            '0.0000000000001',
            True,
        ):
            with self.subTest(value=value), self.assertRaises(ValueError):
                normalize_baseline(dict(self.ratio(), target_margin_pct=value))
        for updates in (
            {'business_order_id': ''},
            {'tax_basis': 'bad'},
            {'target_defect_pct': '101'},
            {'quantity': '0'},
            {'unit_price': '-1'},
            {'discount_per_pair': '11'},
        ):
            with self.subTest(updates=updates), self.assertRaises(ValueError):
                normalize_baseline(self.costs(**updates))

    def test_excel_key_value_percent_formats(self):
        for value, fmt in [(20, None), ('20%', None), (0.2, 'B4')]:
            content = workbook(
                [
                    ['字段', '数值'],
                    ['计算方式', 'ratio'],
                    ['业务订单号', '业务001'],
                    ['目标生产毛利率', value],
                ],
                fmt,
            )
            result = preview_theory(content)
            self.assertEqual(result['baseline']['target_margin_pct'], '20')
            self.assertEqual(result['source_hash'], hashlib.sha256(content).hexdigest())

    def test_excel_header_layout_and_duplicate_rejection(self):
        result = preview_theory(
            workbook(
                [
                    ['method', 'business_order_id', 'target_margin_pct'],
                    ['ratio', 'A', '20%'],
                ]
            )
        )
        self.assertEqual(result['baseline']['business_order_id'], 'A')
        with self.assertRaisesRegex(ValueError, '重复'):
            preview_theory(
                workbook(
                    [
                        ['字段', '数值'],
                        ['method', 'ratio'],
                        ['business_order_id', 'A'],
                        ['target_margin_pct', 20],
                        ['目标生产毛利率', 20],
                    ]
                )
            )

    def test_excel_rejects_formulas_unknown_layout_and_limits(self):
        for rows in (
            [['字段', '数值'], ['method', '=1+1']],
            [['随便标题'], ['A', 'B']],
            [
                ['method', 'business_order_id', 'target_margin_pct'],
                ['ratio', 'A', 20],
                ['ratio', 'B', 20],
            ],
        ):
            with self.subTest(rows=rows), self.assertRaises(ValueError):
                preview_theory(workbook(rows))
        content = workbook(
            [
                ['字段', '数值'],
                ['method', 'ratio'],
                ['business_order_id', 'A'],
                ['target_margin_pct', 20],
            ]
        )
        book = load_workbook(io.BytesIO(content))
        book.create_sheet('hidden')['A1'] = '=1+1'
        stream = io.BytesIO()
        book.save(stream)
        with self.assertRaisesRegex(ValueError, '公式'):
            preview_theory(stream.getvalue())
        with self.assertRaises(ValueError):
            preview_theory(b'0' * (10 * 1024 * 1024 + 1))
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('oversize', b'0' * (50 * 1024 * 1024 + 1))
        with self.assertRaises(ValueError):
            preview_theory(stream.getvalue())

    def test_template_leaves_business_values_blank(self):
        book = load_workbook(io.BytesIO(theory_template()))
        self.assertIn('理论利润表', book.sheetnames)
        values = {
            row[0]: row[1]
            for row in book['理论利润表'].iter_rows(min_row=2, values_only=True)
        }
        self.assertIsNone(values['目标生产毛利率'])
        self.assertIsNone(values['理论材料成本'])
        with self.assertRaises(ValueError):
            preview_theory(theory_template())

    def test_excel_dimension_cannot_hide_formula_or_out_of_range_cells(self):
        content = workbook(
            [
                ['字段', '数值'],
                ['method', 'ratio'],
                ['business_order_id', 'A'],
                ['target_margin_pct', 20],
            ]
        )
        for extra, error in [
            ('<row r="5"><c r="A5"><f>1+1</f><v>2</v></c></row>', '公式'),
            ('<row r="1001"><c r="A1001"><v>2</v></c></row>', '范围'),
        ]:
            stream = io.BytesIO()
            with (
                zipfile.ZipFile(io.BytesIO(content)) as source,
                zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as target,
            ):
                for name in source.namelist():
                    data = source.read(name)
                    if name == 'xl/worksheets/sheet1.xml':
                        data = data.replace(
                            b'</sheetData>', extra.encode() + b'</sheetData>'
                        )
                    target.writestr(name, data)
            with self.subTest(error=error), self.assertRaisesRegex(ValueError, error):
                preview_theory(stream.getvalue())


if __name__ == '__main__':
    unittest.main()
