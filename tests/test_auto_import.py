import io
import unittest
from openpyxl import load_workbook
from test_eva_source import source_bytes
from profit_model.excel_import import FIELDS,preview_excel


def labeled_source(change=None):
    w=load_workbook(io.BytesIO(source_bytes()))
    for name in ('原订单用量','补单用量'):
        s=w[name]
        for cell,value in [('B5','材料编号'),('C5','材料成本'),('D5','配方用量')]:s[cell]=value
        for row in (5,9,13):s[f'F{row}']='码号';s[f'F{row+1}']='SOP重量换算';s[f'F{row+2}']='订单数量'
    s=w['生产周期']
    for cell,value in [('A5','生产日期'),('B5','生产数量'),('F5','不良数'),('J5','生产欠数')]:s[cell]=value
    for first,second,third in [('B','F','J'),('C','G','K'),('D','H','L')]:s[second+'6']=s[first+'6'].value;s[third+'6']=s[first+'6'].value
    s=w['成本计算'];s['B12']='加工成本';s['B26']='行政成本'
    for fid,label,_,_ in FIELDS:s['C'+fid[1:]]=label;s['F'+fid[1:]]='双' if fid=='E8' else '元/双'
    if change:change(w)
    out=io.BytesIO();w.save(out);return out.getvalue()


class AutoImportTests(unittest.TestCase):
    def test_clear_layout_auto_adapts_moved_fee_without_manual_confirmation(self):
        from profit_model.excel_import import auto_excel
        def change(w):
            s=w['成本计算'];s['C17']=None;s['E17']=999;s['C40']='生产人员工资';s['E40']=2;s['F40']='元/双'
        outcome=auto_excel(labeled_source(change))
        self.assertEqual(outcome['status'],'ready')
        a=outcome['data']['source_analysis']
        self.assertEqual(a['metrics']['production_profit'],'1562.50')
        self.assertEqual(a['import_review']['mode'],'automatic')
        self.assertFalse(a['import_review']['confirmed'])

    def test_missing_unit_ambiguous_field_and_unknown_layout_require_review(self):
        from profit_model.excel_import import auto_excel
        for change in [lambda w:w['成本计算'].__setitem__('F17',None),lambda w:w['原订单用量'].__setitem__('D5','金额'),lambda w:w['成本计算'].__setitem__('E17',None)]:
            with self.subTest(change=change):
                outcome=auto_excel(labeled_source(change))
                self.assertEqual(outcome['status'],'needs_review');self.assertNotIn('data',outcome)
        def duplicate(w):w['成本计算']['C40']='直接人工';w['成本计算']['E40']=2
        self.assertEqual(auto_excel(labeled_source(duplicate))['status'],'needs_review')

    def test_profile_reuse_requires_matching_structure_and_approval(self):
        from profit_model.excel_import import auto_excel
        content=source_bytes();p=preview_excel(content)
        profile={'name':'已确认模板','approved':True,'layout_signature':p['layout_signature'],'mapping':p['mapping']}
        out=auto_excel(content,[profile])
        self.assertEqual(out['status'],'ready');self.assertEqual(out['mode'],'automatic_profile')
        self.assertEqual(auto_excel(content,[dict(profile,approved=False)])['status'],'needs_review')
        w=load_workbook(io.BytesIO(content));w['成本计算']['F17']='元/小时';b=io.BytesIO();w.save(b)
        self.assertEqual(auto_excel(b.getvalue(),[profile])['status'],'needs_review')

    def test_production_outside_supported_rows_is_not_silently_omitted(self):
        from profit_model.excel_import import auto_excel
        def change(w):
            s=w['生产周期'];s['A44']=46238;s['B44']=1;s['F44']=0
        self.assertEqual(auto_excel(labeled_source(change))['status'],'needs_review')

    def test_formula_material_and_text_or_formula_daily_overflow_require_review(self):
        from profit_model.excel_import import auto_excel
        def material(w):
            s=w['原订单用量'];s['B21']='M3';s['C21']='=10';s['D21']='=1'
        changes=[material]
        for value in ('2026-10-14','=46238'):
            def daily(w,value=value):
                s=w['生产周期'];s['A44']=value;s['B44']=1
            changes.append(daily)
        for change in changes:
            self.assertEqual(auto_excel(labeled_source(change))['status'],'needs_review')
