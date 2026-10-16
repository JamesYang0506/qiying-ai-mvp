import base64
import io
import json
import unittest
import urllib.error
import urllib.request
from pathlib import Path

from openpyxl import Workbook
from test_server import ServerTests as _Setup
from test_engine import fixture
from test_comparison import LEGACY_RULES, RELATIVE_RULES


class ComparisonAPITests(unittest.TestCase):
    setUp = _Setup.setUp
    tearDown = _Setup.tearDown

    def post(self, path, payload):
        request = urllib.request.Request(
            self.url + path,
            data=json.dumps(payload).encode(),
            headers={'Content-Type': 'application/json'},
        )
        with urllib.request.urlopen(request) as response:
            return json.load(response)

    def import_order(self):
        return self.post('/api/calculate', fixture())

    def preview(self, workbook, selection=None):
        output = io.BytesIO()
        workbook.save(output)
        payload = {'workbook_base64': base64.b64encode(output.getvalue()).decode()}
        if selection is not None:
            payload['selection'] = selection
        return self.post('/api/theory/preview', payload)

    def stored_files(self):
        root = Path(self.tmp.name)
        return {
            str(path.relative_to(root)): path.read_bytes()
            for path in root.rglob('*')
            if path.is_file()
        }

    def save_existing_baseline(self):
        order = self.import_order()
        self.post(
            '/api/theory',
            {
                'snapshot_id': order['snapshot']['snapshot_id'],
                'order_id': 'A',
                'confirmed': True,
                'scope_confirmed': True,
                'baseline': {
                    'method': 'ratio',
                    'business_order_id': 'A',
                    'target_margin_pct': '40',
                    'quantity': '1000',
                    'unit': '双',
                },
            },
        )

    def test_enterprise_offset_preview_and_partial_unknown_preserve_storage(self):
        self.save_existing_baseline()
        before = self.stored_files()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '客户计划'
        sheet['B2'] = '客户报价计划'
        for row, label, value in [(4, '订单编号', 'A'), (5, '计划毛利率', '25%')]:
            sheet.cell(row, 2, label)
            sheet.cell(row, 3, value)
        result = self.preview(workbook)
        self.assertEqual(result['status'], 'ready')
        self.assertEqual(result['baseline']['business_order_id'], 'A')
        self.assertEqual(result['baseline']['target_margin_pct'], '25')
        source = next(
            field
            for field in result['adaptation']['fields']
            if field['field'] == 'target_margin_pct'
        )
        self.assertEqual(source['cell'], 'C5')
        self.assertEqual(result['adaptation']['sheet'], '客户计划')
        self.assertEqual(len(result['source_hash']), 64)
        self.assertIsInstance(result['warnings'], list)
        sheet['B5'] = '企业自定义备注'
        sheet['C5'] = '待人工确认目标'
        partial = self.preview(workbook)
        self.assertEqual(partial['status'], 'needs_review')
        self.assertEqual(partial['baseline']['business_order_id'], 'A')
        self.assertTrue(partial['adaptation']['missing'])
        self.assertEqual(self.stored_files(), before)

    def test_sheet_and_record_selection_preview_remains_readonly(self):
        self.save_existing_baseline()
        before = self.stored_files()
        workbook = Workbook()
        workbook.active.title = '说明'
        workbook.active['A1'] = '上传预算后请核对来源'
        for name in ('预算一', '预算二'):
            sheet = workbook.create_sheet(name)
            for column, value in enumerate(('业务订单号', '目标生产毛利率'), 2):
                sheet.cell(4, column, value)
            for row, order_id, margin in [(5, 'A', '25%'), (6, 'B', '30%')]:
                sheet.cell(row, 2, order_id)
                sheet.cell(row, 3, margin)
        unresolved = self.preview(workbook)
        self.assertEqual(unresolved['status'], 'needs_review')
        selected = self.preview(workbook, {'sheet': '预算二', 'record_row': 6})
        self.assertEqual(selected['status'], 'ready')
        self.assertEqual(selected['baseline']['business_order_id'], 'B')
        self.assertEqual(selected['baseline']['target_margin_pct'], '30')
        self.assertEqual(selected['adaptation']['sheet'], '预算二')
        self.assertEqual(self.stored_files(), before)

    def test_invalid_selectors_base64_and_upload_size_are_rejected(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['业务订单号', '目标生产毛利率'])
        sheet.append(['A', '25%'])
        for selection in (
            'bad',
            {'sheet': '不存在'},
            {'record_row': -1},
            {'mapping': {'target_margin_pct': 'Other!B2'}},
        ):
            with self.subTest(selection=selection):
                with self.assertRaises(urllib.error.HTTPError) as error:
                    self.preview(workbook, selection)
                self.assertEqual(error.exception.code, 400)
        for encoded in (
            'not-base64!',
            base64.b64encode(b'x' * (10 * 1024 * 1024 + 1)).decode(),
        ):
            with self.subTest(encoded_length=len(encoded)):
                with self.assertRaises(urllib.error.HTTPError) as error:
                    self.post('/api/theory/preview', {'workbook_base64': encoded})
                self.assertEqual(error.exception.code, 400)
        self.assertEqual(self.stored_files(), {})

    def test_import_auto_checks_and_baseline_rules_history_remain_separate(self):
        self.post('/api/comparison/rules', {'confirmed': True, 'rules': LEGACY_RULES})
        order = self.import_order()
        sid = order['snapshot']['snapshot_id']
        with urllib.request.urlopen(
            self.url + '/api/comparison?snapshot_id=' + sid + '&order_id=A'
        ) as response:
            workspace = json.load(response)
        self.assertEqual(
            workspace['comparison']['metrics']['margin']['status'], 'pending'
        )
        result = self.post(
            '/api/theory',
            {
                'snapshot_id': sid,
                'order_id': 'A',
                'confirmed': True,
                'scope_confirmed': True,
                'baseline': {
                    'method': 'ratio',
                    'business_order_id': 'A',
                    'target_margin_pct': '40',
                    'quantity': '1000',
                    'unit': '双',
                },
            },
        )
        self.assertEqual(result['comparison']['metrics']['margin']['status'], 'danger')
        self.assertEqual(result['comparison']['rules']['margin_warning_pp'], '3')
        rules = self.post(
            '/api/comparison/rules',
            {
                'confirmed': True,
                'rules': {
                    'margin_warning_pp': '10',
                    'margin_danger_pp': '15',
                    'defect_warning_pct': '3',
                    'defect_danger_pct': '5',
                    'min_production': '100',
                },
            },
        )
        self.assertNotEqual(
            rules['rules']['version'], result['comparison']['rules']['version']
        )
        updated = self.post(
            '/api/comparison/check', {'snapshot_id': sid, 'order_id': 'A'}
        )
        self.assertEqual(updated['comparison']['metrics']['margin']['status'], 'normal')
        self.assertEqual(len(updated['records']), 3)

    def test_relative_rules_roundtrip_and_check_preserves_prior_record(self):
        order = self.import_order()
        sid = order['snapshot']['snapshot_id']
        legacy = self.post(
            '/api/comparison/rules', {'confirmed': True, 'rules': LEGACY_RULES}
        )
        before = self.post(
            '/api/theory',
            {
                'snapshot_id': sid,
                'order_id': 'A',
                'confirmed': True,
                'scope_confirmed': True,
                'baseline': {
                    'method': 'ratio',
                    'business_order_id': 'A',
                    'target_margin_pct': '40',
                    'quantity': '1000',
                    'unit': '双',
                },
            },
        )
        saved = self.post(
            '/api/comparison/rules', {'confirmed': True, 'rules': RELATIVE_RULES}
        )
        self.assertEqual(saved['rules']['margin_mode'], 'relative')
        self.assertNotEqual(saved['rules']['version'], legacy['rules']['version'])
        result = self.post(
            '/api/comparison/check', {'snapshot_id': sid, 'order_id': 'A'}
        )
        margin = result['comparison']['metrics']['margin']
        self.assertEqual(margin['status'], 'focus')
        self.assertIsNotNone(margin['relative_drop_pct'])
        frozen = next(
            row for row in result['records'] if row['id'] == before['comparison']['id']
        )
        self.assertEqual(frozen, before['comparison'])

    def test_template_preview_and_static_assets(self):
        for path in (
            '/comparison.js',
            '/theory-import.js',
            '/comparison.css',
            '/api/theory/template.xlsx',
        ):
            with urllib.request.urlopen(self.url + path) as response:
                self.assertTrue(response.read())
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '理论利润表'
        for row in [
            ('字段', '数值'),
            ('method', 'ratio'),
            ('business_order_id', 'A'),
            ('target_margin_pct', '20%'),
        ]:
            sheet.append(row)
        output = io.BytesIO()
        workbook.save(output)
        result = self.post(
            '/api/theory/preview',
            {'workbook_base64': base64.b64encode(output.getvalue()).decode()},
        )
        self.assertEqual(result['baseline']['target_margin_pct'], '20')

    def test_unknown_snapshot_deleted_records_and_unconfirmed_writes_rejected(self):
        order = self.import_order()
        sid = order['snapshot']['snapshot_id']
        self.post(
            '/api/history/delete',
            {'snapshot_id': sid, 'order_id': 'A', 'confirmed': True},
        )
        for path, payload in [
            ('/api/comparison/check', {'snapshot_id': sid, 'order_id': 'A'}),
            ('/api/comparison/rules', {'confirmed': False, 'rules': {}}),
            (
                '/api/theory/bind',
                {
                    'confirmed': True,
                    'snapshot_id': '../latest',
                    'order_id': 'A',
                    'baseline_id': '../bad',
                },
            ),
        ]:
            with (
                self.subTest(path=path),
                self.assertRaises(urllib.error.HTTPError) as error,
            ):
                self.post(path, payload)
            self.assertEqual(error.exception.code, 400)


del _Setup
