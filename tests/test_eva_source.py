import io
import unittest
from openpyxl import Workbook,load_workbook
from profit_model.eva_source import FormulaReader, read_eva_workbook
from profit_model.engine import calculate


def source_bytes():
    w=Workbook();w.remove(w.active)
    for name in ['原订单用量','补单用量','生产周期','成本计算']:w.create_sheet(name)
    for name,qty in [('原订单用量',100),('补单用量',10)]:
        s=w[name];s['B3']='TEST';s['E3']='黑';s['H3']=qty
        s['H5']='4.5#'
        for row,code,price,weight in [(6,'M1',10,3),(7,'M2',20,1)]:
            s[f'B{row}']=code;s[f'C{row}']=price;s[f'D{row}']=weight;s[f'E{row}']=f'=C{row}*D{row}'
        s['H6']=.1;s['H7']=qty;s['H8']='=H6*H7'
        s['K3']='=SUM(G8:K8)+SUM(G12:K12)+SUM(G16:K16)'
        s['D18']='=SUM(D6:D17)';s['E18']='=SUM(E6:E17)';s['B20']='=E18/D18';s['B19']='=B20*K3'
    s=w['生产周期'];s['B3']='TEST';s['F3']=100;s['J7']=100;s['K7']=0;s['L7']=0;s['A8']=46237;s['B6']='4.5#';s['C6']='8#';s['D6']='9.5#'
    for c,v in [('B',108),('C',0),('D',0),('F',8),('G',0),('H',0)]:
        s[f'{c}8']=v;s[f'{c}44']=f'=SUM({c}8:{c}43)'
    s=w['成本计算'];s['A1']='订单成本 - TEST';s['E4']=20;s['E5']='=E4*0.05';s['H5']='=E4-E5';s['E8']=100
    s['E9']='=原订单用量!B19';s['E10']='=补单用量!B19';s['H10']='=(E9+E10)/E8'
    for row in range(13,24):s[f'C{row}']='费用'+str(row);s[f'E{row}']=1 if row==17 else 0
    for row in range(27,33):s[f'C{row}']='行政'+str(row);s[f'E{row}']=.2 if row==27 else 0
    s['H23']='=SUM(E13:E23)';s['H24']='=H5-H10-H23';s['H32']='=SUM(E27:E32)';s['H33']='=H24-H32'
    b=io.BytesIO();w.save(b);return b.getvalue()


class SourceTests(unittest.TestCase):
    def test_remaining_discrepancy_is_detected_without_changing_costs(self):
        baseline=read_eva_workbook(source_bytes())
        w=load_workbook(io.BytesIO(source_bytes()))
        for col in ('J','K','L'):w['生产周期'][col+'8']=0
        w['生产周期']['J8']=1
        b=io.BytesIO();w.save(b)
        a=read_eva_workbook(b.getvalue())['source_analysis']
        check=next(x for x in a['trust_review']['checks'] if x['code']=='REMAINING')
        self.assertEqual(check['status'],'difference')
        self.assertIn('J8',check['evidence'])
        self.assertEqual(a['metrics'],baseline['source_analysis']['metrics'])

    def test_missing_remaining_records_are_not_a_false_pass(self):
        a=read_eva_workbook(source_bytes())['source_analysis']
        self.assertEqual(next(x['status'] for x in a['trust_review']['checks'] if x['code']=='REMAINING'),'pending')

    def test_review_groups_and_prioritizes_issues_without_false_quality_pass(self):
        review=read_eva_workbook(source_bytes())['source_analysis']['trust_review']
        self.assertEqual({x['id'] for x in review['categories']},{'profit','consistency','quality','process'})
        self.assertEqual(sum(sum(x['counts'].values()) for x in review['categories']),len(review['checks']))
        priority={'difference':0,'pending':1,'passed':2}
        order=[priority[x['status']] for x in review['checks']]
        self.assertEqual(order,sorted(order))
        for code in ('QUALITY_RATE','PROCESS_PARAMETERS'):
            item=next(x for x in review['checks'] if x['code']==code)
            self.assertEqual(item['status'],'pending')
            self.assertEqual(item['judgment'],'暂无法判断')

    def test_two_profit_layers_and_review(self):
        a=read_eva_workbook(source_bytes())['source_analysis'];v=a['profit_layers']
        self.assertEqual(v['production_cost'],'237.50')
        self.assertEqual(v['operating_expense'],'20.00')
        self.assertEqual(v['cumulative_cost'],'257.50')
        self.assertEqual(v['operating_net_profit'],'1642.50')
        statuses={x['code']:x['status'] for x in a['trust_review']['checks']}
        self.assertEqual(statuses['QUANTITY'],'passed')
        self.assertEqual(statuses['REPLENISHMENT'],'difference')
        self.assertEqual(statuses['NET_SCOPE'],'pending')

    def test_missing_operating_cost_blocks_net_layer(self):
        w=load_workbook(io.BytesIO(source_bytes()));w['成本计算']['E27']=None
        b=io.BytesIO();w.save(b);a=read_eva_workbook(b.getvalue())['source_analysis']
        self.assertEqual(a['profit_layers']['production_cost'],'237.50')
        self.assertIsNone(a['profit_layers']['operating_net_profit'])
        self.assertIsNone(a['profit_layers']['operating_expense'])
        self.assertIsNone(a['profit_layers']['management_expense'])
        self.assertEqual(next(x['status'] for x in a['trust_review']['checks'] if x['code']=='REQUIRED_RATES'),'difference')

    def test_profit_and_subtotals_are_not_calculation_inputs(self):
        baseline=read_eva_workbook(source_bytes())['source_analysis']['metrics']
        for replacement in (999999, None, '=H33+1'):
            w=load_workbook(io.BytesIO(source_bytes()))
            for cell in ('H5','H10','E9','E10','H23','H24','H32','H33'):
                w['成本计算'][cell]=replacement
            for sheet in ('原订单用量','补单用量'):
                for cell in ('B19','B20','K3','D18','E18'):
                    w[sheet][cell]=replacement
            b=io.BytesIO();w.save(b)
            actual=read_eva_workbook(b.getvalue())['source_analysis']['metrics']
            self.assertEqual(actual,baseline)

    def test_input_cost_changes_profit_without_using_answer(self):
        w=load_workbook(io.BytesIO(source_bytes()));w['成本计算']['E17']=2
        b=io.BytesIO();w.save(b)
        m=read_eva_workbook(b.getvalue())['source_analysis']['metrics']
        self.assertEqual(m['production_profit'],'1562.50')
        self.assertEqual(m['management_reference_profit'],'1542.50')

    def test_input_cannot_indirectly_reference_profit(self):
        w=load_workbook(io.BytesIO(source_bytes()));w['成本计算']['E17']='=H24'
        b=io.BytesIO();w.save(b)
        with self.assertRaisesRegex(ValueError,'不能作为独立计算输入'):
            read_eva_workbook(b.getvalue())

    def test_local_formula_recomputed_not_trusted_cache(self):
        w=Workbook();c=Workbook();w.active.title=c.active.title='S'
        w['S']['A1']=10;w['S']['A2']=20;w['S']['B1']='=SUM(A1:A2)*0.5';c['S']['B1']=999
        r=FormulaReader(w,c)
        self.assertEqual(str(r.value('S','B1')),'15.0')

    def test_external_formula_uses_explicit_cache(self):
        w=Workbook();c=Workbook();w.active.title=c.active.title='S'
        w['S']['A1']="='[1]outside'!M22";c['S']['A1']=.25
        r=FormulaReader(w,c);self.assertEqual(str(r.value('S','A1')),'0.25')
        self.assertEqual(len(r.external),1)
        c['S']['A1']=None
        self.assertIsNone(FormulaReader(w,c).value('S','A1'))

    def test_source_profit_and_defect_reconciliation(self):
        d=read_eva_workbook(source_bytes());a=d['source_analysis']
        self.assertEqual(a['metrics']['revenue'],'1900.00')
        self.assertEqual(a['metrics']['material_total'],'137.50')
        self.assertEqual(a['metrics']['production_profit'],'1662.50')
        self.assertEqual(a['metrics']['management_reference_profit'],'1642.50')
        self.assertEqual(a['production']['defect_quantity'],'8')
        self.assertEqual(a['production']['replenishment_quantity'],'10')
        self.assertEqual(a['production']['difference'],'-2')
        r=calculate(d)
        self.assertIsNone(r['orders'][0]['production_profit'])
        self.assertIsNone(r['orders'][0]['management_profit'])
        self.assertEqual(r['source_analysis']['metrics']['management_reference_profit'],'1642.50')

    def test_source_layout_invalid_rejected(self):
        with self.assertRaises(ValueError):read_eva_workbook(b'invalid')

    def test_missing_source_wage_not_interpreted_as_zero(self):
        w=load_workbook(io.BytesIO(source_bytes()));w['成本计算']['E17']=None
        b=io.BytesIO();w.save(b);d=read_eva_workbook(b.getvalue())
        self.assertIsNone(d['source_analysis']['metrics']['processing_total'])
        self.assertIsNone(d['source_analysis']['metrics']['production_profit'])
        self.assertEqual(next(x for x in d['costs'] if x['record_id']=='source-E17')['status'],'missing')

    def test_negative_production_size_cannot_be_hidden_by_total(self):
        w=load_workbook(io.BytesIO(source_bytes()));w['生产周期']['B8']=-1;w['生产周期']['C8']=109
        b=io.BytesIO();w.save(b)
        with self.assertRaises(ValueError):read_eva_workbook(b.getvalue())
