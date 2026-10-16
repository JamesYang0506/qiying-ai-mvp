import io
import unittest
import zipfile

from openpyxl import Workbook

from profit_model.theory_import import auto_preview_theory


def workbook(rows, other=None, formats=None):
    book = Workbook()
    sheet = book.active
    sheet.title = '核价'
    for row in rows:
        sheet.append(row)
    for coordinate, fmt in (formats or {}).items():
        sheet[coordinate].number_format = fmt
    if other:
        second = book.create_sheet('第二表')
        for row in other:
            second.append(row)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


class AutoTheoryTests(unittest.TestCase):
    def test_budget_aliases_and_blank_labor_never_use_derived_neighbor(self):
        rows = [
            ['订单号', '001'],
            ['订单数量', 100],
            ['工厂报价', 17.2],
            ['BOM材料', '65%'],
            ['生产管理工资', '2.8%'],
            ['生产损耗预算', '0.6%'],
            ['不良品损耗预算', '1.8%'],
            ['电汽动力能源', '0.4%'],
            ['备料计件人工（元/双）', None, '=0'],
            ['油压计件人工（元/双）', None, '=0'],
            ['公司管理干部', '1.3%'],
            ['行政人员', '1.7%'],
            ['扣计件人工前毛利率', '29.4%'],
        ]
        preview = auto_preview_theory(workbook(rows))
        self.assertEqual(preview['baseline']['method'], 'quote')
        self.assertEqual(preview['baseline'].get('production_loss_pct'), '0.6')
        self.assertEqual(preview['baseline'].get('defect_loss_pct'), '1.8')
        self.assertEqual(preview['baseline'].get('energy_pct'), '0.4')
        self.assertEqual(
            preview['adaptation']['missing'],
            ['preparation_labor_per_pair', 'press_labor_per_pair'],
        )
        self.assertNotIn('target_margin_pct', preview['baseline'])
        source = next(
            f
            for f in preview['adaptation']['fields']
            if f['field'] == 'preparation_labor_per_pair'
        )
        self.assertEqual(source['cell'], 'B9')
        margin = auto_preview_theory(
            workbook([['订单号', '001'], ['完整计划生产毛利率', '20%']])
        )
        self.assertEqual(margin['baseline']['target_margin_pct'], '20')

    def test_unit_column_headings_are_not_empty_unit_conflicts(self):
        preview = auto_preview_theory(
            workbook(
                [
                    ['说明', '数值', '单位'],
                    ['订单号', '001'],
                    ['完整计划毛利率', '待补人工'],
                    ['说明2', '数值', '单位'],
                ]
            )
        )
        self.assertIsNone(preview['baseline'].get('target_margin_pct'))
        self.assertIn('target_margin_pct', preview['adaptation']['missing'])
        self.assertFalse(
            any('多个候选' in note for note in preview['adaptation']['notes'])
        )

    def test_generic_material_cost_percent_becomes_ratio_not_money(self):
        for value, formats in [('65%', {}), (0.65, {'B2': '0.0%'})]:
            preview = auto_preview_theory(
                workbook([['订单号', '001'], ['材料成本', value]], formats=formats)
            )
            self.assertEqual(preview['baseline'].get('material_pct'), '65')
            self.assertNotIn('material_cost', preview['baseline'])
            self.assertEqual(preview['baseline']['method'], 'quote')

    def test_explicit_cost_method_conflicting_material_percentage_not_ready(self):
        preview = auto_preview_theory(
            workbook(
                [
                    ['method', 'costs'],
                    ['订单号', '001'],
                    ['理论材料成本%', 50],
                ]
            )
        )
        self.assertEqual(preview['status'], 'needs_review')
        self.assertEqual(preview['baseline']['material_pct'], 50)
        self.assertNotIn('material_cost', preview['baseline'])
        self.assertEqual(
            preview['adaptation']['conflicts'][0]['code'], 'mixed_cost_basis'
        )

    def test_units_are_validated_for_each_field_including_excel_formats(self):
        for label, field, value, formats in [
            ('能源', 'energy_pct', '0.4元/双', {}),
            ('能源', 'energy_pct', '0.4元', {}),
            ('能源', 'energy_pct', '0.4双', {}),
            ('工厂报价', 'unit_price', '17.2双', {}),
            ('数量', 'quantity', '100元', {}),
            ('人工成本', 'labor_cost', '65%', {}),
            ('加工制造成本', 'manufacturing_cost', 0.65, {'B2': '0%'}),
            ('工厂报价', 'unit_price', 0.65, {'B2': '0%'}),
            ('能源', 'energy_pct', 0.4, {'B2': '0.0"元/双"'}),
        ]:
            with self.subTest(label=label, value=value, formats=formats):
                preview = auto_preview_theory(
                    workbook([['订单号', '001'], [label, value]], formats=formats)
                )
                self.assertIsNone(preview['baseline'].get(field))
                self.assertTrue(
                    any('单位' in warning for warning in preview['warnings'])
                )

    def test_offset_headers_records_require_selection(self):
        content = workbook(
            [
                ['标题'],
                [],
                [None, '订单号', '目标毛利率'],
                [None, '0012', '20%'],
                [None, '0013', '25%'],
            ]
        )
        preview = auto_preview_theory(content)
        self.assertEqual(preview['status'], 'needs_review')
        self.assertEqual(preview['baseline'], {})
        self.assertEqual([r['row'] for r in preview['adaptation']['records']], [4, 5])
        preview = auto_preview_theory(content, {'record_row': 4})
        self.assertEqual(preview['status'], 'ready')
        self.assertEqual(preview['baseline']['business_order_id'], '0012')
        self.assertEqual(preview['baseline']['target_margin_pct'], '20')

    def test_quote_aliases_zero_and_excel_percent(self):
        rows = [
            ['虚拟报价'],
            ['定单号', '0012'],
            ['数量', 100],
            ['工厂报价（元/双）', 17.2],
            ['BOM材料', 0.65],
            ['生产管理工资', '2.8%'],
            ['生产损耗', '0.6%'],
            ['不良品损耗', '1.8%'],
            ['能源', '0.4%'],
            ['备料人工（元/双）', 0],
            ['油压人工（元/双）', 0],
            ['公司管理干部', '1.3%'],
            ['行政', '1.7%'],
        ]
        preview = auto_preview_theory(workbook(rows, formats={'B5': '0%'}))
        self.assertEqual(preview['status'], 'ready')
        self.assertEqual(preview['baseline']['method'], 'quote')
        self.assertEqual(preview['baseline']['material_pct'], '65')
        self.assertEqual(preview['baseline']['preparation_labor_per_pair'], '0')
        self.assertIsNone(preview['baseline']['target_defect_pct'])

    def test_partial_and_ambiguous_candidates_not_guessed(self):
        content = workbook([['订单号', '001'], ['目标毛利率', 20], ['目标毛利率', 30]])
        result = auto_preview_theory(content)
        self.assertEqual(result['status'], 'needs_review')
        self.assertIn('target_margin_pct', result['adaptation']['missing'])
        self.assertIsNone(result['baseline'].get('target_margin_pct'))
        field = next(
            f
            for f in result['adaptation']['fields']
            if f['field'] == 'target_margin_pct'
        )
        self.assertEqual(len(field['candidates']), 2)
        chosen = auto_preview_theory(content, {'mapping': {'target_margin_pct': 'B3'}})
        self.assertEqual(chosen['baseline']['target_margin_pct'], '30')

    def test_mapping_preserves_source_unit_validation(self):
        content = workbook(
            [
                ['订单号', '001'],
                ['目标毛利率', 20],
                ['材料成本（元/双）', 12],
                ['材料成本', 50],
            ]
        )
        result = auto_preview_theory(content, {'mapping': {'material_cost': 'B3'}})
        self.assertIsNone(result['baseline'].get('material_cost'))
        field = next(
            f for f in result['adaptation']['fields'] if f['field'] == 'material_cost'
        )
        self.assertIsNone(field['value'])
        with self.assertRaises(ValueError):
            auto_preview_theory(content, {'mapping': {'material_cost': 'B2'}})

    def test_multiple_sheets_and_invalid_selection(self):
        rows = [['订单号', '001'], ['目标毛利率', 20]]
        content = workbook(rows, other=rows)
        self.assertEqual(auto_preview_theory(content)['baseline'], {})
        self.assertEqual(
            auto_preview_theory(content, {'sheet': '第二表'})['status'], 'ready'
        )
        for selection in (
            {'sheet': 'missing'},
            {'sheet': '核价', 'mapping': {'quantity': 'A999'}},
            {'sheet': '核价', 'mapping': {'quantity': "第二表!B1"}},
        ):
            with self.subTest(selection=selection), self.assertRaises(ValueError):
                auto_preview_theory(content, selection)

    def test_sheet_with_most_fields_is_auto_chosen_without_ties(self):
        # 计划表字段多、对照页只剩一个字段：不再停在"请选工作表"。
        preview = auto_preview_theory(
            workbook(
                [
                    ['订单号', '001'],
                    ['订单数量', 100],
                    ['目标毛利率', 20],
                ],
                other=[['完整计划毛利率', 20]],
            )
        )
        self.assertEqual(preview['adaptation']['sheet'], '核价')
        self.assertEqual(preview['baseline']['business_order_id'], '001')
        self.assertEqual(
            [row['name'] for row in preview['adaptation']['sheets']],
            ['核价', '第二表'],
        )
        self.assertTrue(
            any('自动选择' in note for note in preview['adaptation']['notes'])
        )

    def test_per_pair_cost_cannot_be_total_cost_and_customer_is_not_order(self):
        result = auto_preview_theory(
            workbook(
                [['下单客户', '客户001'], ['材料成本（元/双）', 12], ['目标毛利率', 20]]
            )
        )
        self.assertIn('business_order_id', result['adaptation']['missing'])
        self.assertIsNone(result['baseline'].get('material_cost'))

    def test_numeric_id_requires_review(self):
        result = auto_preview_theory(workbook([['订单号', 123], ['目标毛利率', 20]]))
        self.assertEqual(result['status'], 'needs_review')
        self.assertTrue(any('订单号' in w and '数字' in w for w in result['warnings']))

    def test_formula_without_cache_missing_cached_formula_warns_external_rejected(self):
        raw = workbook([['订单号', '001'], ['目标毛利率', '=10+10']])
        preview = auto_preview_theory(raw)
        self.assertIn('target_margin_pct', preview['adaptation']['missing'])
        source = zipfile.ZipFile(io.BytesIO(raw))
        target = io.BytesIO()
        with zipfile.ZipFile(target, 'w') as out:
            for member in source.infolist():
                data = source.read(member.filename)
                if member.filename == 'xl/worksheets/sheet1.xml':
                    data = data.replace(b'<v></v>', b'<v>20</v>')
                out.writestr(member, data)
        cached = auto_preview_theory(target.getvalue())
        self.assertEqual(cached['baseline']['target_margin_pct'], '20')
        self.assertTrue(any('缓存' in w for w in cached['warnings']))
        external = auto_preview_theory(
            workbook([['订单号', '001'], ['目标毛利率', "='[other.xlsx]Sheet1'!A1"]])
        )
        self.assertIn('target_margin_pct', external['adaptation']['missing'])

    def test_security_limits_and_empty(self):
        for content in (b'', b'not excel', b'x' * (10 * 1024 * 1024 + 1)):
            with self.assertRaises(ValueError):
                auto_preview_theory(content)
        with self.assertRaises(ValueError):
            auto_preview_theory(workbook([[]] * 1000 + [['订单号', '001']]))
        with self.assertRaises(ValueError):
            auto_preview_theory(workbook([['nothing']]))

    def test_several_neighboring_pairs_on_same_row(self):
        preview = auto_preview_theory(
            workbook(
                [
                    ['订单编号', '0012', None, '计划毛利率', '25%'],
                ]
            )
        )
        self.assertEqual(preview['status'], 'ready')
        self.assertEqual(preview['baseline']['business_order_id'], '0012')

    def test_costs_are_independent_and_derived_margin_cannot_override(self):
        rows = [
            ['method', 'costs'],
            ['订单号', '001'],
            ['数量', 10],
            ['报价单价', 20],
            ['每双折扣', 0],
            ['材料成本', 50],
            ['人工成本', 30],
            ['加工制造成本', 10],
            ['其他生产成本', 0],
            ['目标毛利率', 99],
        ]
        preview = auto_preview_theory(workbook(rows))
        self.assertEqual(preview['status'], 'needs_review')
        self.assertEqual(preview['baseline']['material_cost'], 50)
        self.assertTrue(any('冲突' in warning for warning in preview['warnings']))

    def test_quote_does_not_overwrite_independent_cost_evidence(self):
        rows = [
            ['method', 'quote'],
            ['订单号', '001'],
            ['数量', 10],
            ['报价单价', 20],
            ['材料成本', 123],
            ['BOM材料', '65%'],
        ]
        preview = auto_preview_theory(workbook(rows))
        self.assertEqual(preview['status'], 'needs_review')
        self.assertEqual(preview['baseline']['material_cost'], 123)
        self.assertTrue(any('独立成本' in warning for warning in preview['warnings']))
        self.assertEqual(
            preview['adaptation']['conflicts'][0]['code'], 'mixed_cost_basis'
        )
        self.assertIn('material_cost', preview['adaptation']['conflicts'][0]['fields'])
        self.assertIn('material_pct', preview['adaptation']['conflicts'][0]['fields'])

    def test_forged_dimensions_and_sheet_count_still_bounded(self):
        raw = workbook([['订单号', '001'], ['目标毛利率', 20]])
        source = zipfile.ZipFile(io.BytesIO(raw))
        target = io.BytesIO()
        with zipfile.ZipFile(target, 'w') as out:
            for member in source.infolist():
                data = source.read(member.filename)
                if member.filename == 'xl/worksheets/sheet1.xml':
                    data = data.replace(b'r="B2"', b'r="B1001"')
                out.writestr(member, data)
        with self.assertRaises(ValueError):
            auto_preview_theory(target.getvalue())
        source = zipfile.ZipFile(io.BytesIO(raw))
        target = io.BytesIO()
        with zipfile.ZipFile(target, 'w') as out:
            for member in source.infolist():
                data = source.read(member.filename)
                if member.filename == 'xl/worksheets/sheet1.xml':
                    data = data.replace(b'ref="A1:B2"', b'ref="A1:B1001"')
                out.writestr(member, data)
        with self.assertRaises(ValueError):
            auto_preview_theory(target.getvalue())
        book = Workbook()
        for i in range(10):
            book.create_sheet(str(i))
        stream = io.BytesIO()
        book.save(stream)
        with self.assertRaises(ValueError):
            auto_preview_theory(stream.getvalue())


if __name__ == '__main__':
    unittest.main()
