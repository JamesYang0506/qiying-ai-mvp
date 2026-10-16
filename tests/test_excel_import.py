import io
import unittest
import base64
import json
import tempfile
import threading
import urllib.request
import urllib.error
from pathlib import Path
from unittest.mock import Mock
from openpyxl import load_workbook
from test_eva_source import source_bytes


def modified(change):
    w=load_workbook(io.BytesIO(source_bytes()));change(w)
    out=io.BytesIO();w.save(out);return out.getvalue()


class ExcelImportTests(unittest.TestCase):
    def test_cost_sheet_alias_and_standard_column_order(self):
        from profit_model.excel_import import preview_excel, confirmed_excel
        from profit_model.files import workbook_bytes
        from profit_model.engine import calculate
        from test_engine import fixture
        content=modified(lambda w:setattr(w['成本计算'],'title','订单成本2026'))
        p=preview_excel(content)
        self.assertEqual(p['mapping']['cost_sheet'],'订单成本2026')
        data=confirmed_excel(content,p['mapping'],p['source_hash'])
        self.assertEqual(data['source_analysis']['metrics']['production_profit'],'1662.50')
        w=load_workbook(io.BytesIO(workbook_bytes(fixture())));s=w['orders']
        for row in s:
            left,right=row[0].value,row[1].value;row[0].value=right;row[1].value=left
        out=io.BytesIO();w.save(out);content=out.getvalue();p=preview_excel(content)
        self.assertEqual(p['kind'],'standard')
        self.assertEqual(calculate(confirmed_excel(content,p['mapping'],p['source_hash']))['orders'][0]['management_profit'],'4350.00')

    def test_numeric_text_is_detected_and_unknown_unit_needs_confirmation(self):
        from profit_model.excel_import import preview_excel
        def change(w):
            s=w['成本计算'];s['C17']=None;s['C40']='生产人员工资';s['E40']='2.00';s['F40']='元/小时'
        p=preview_excel(modified(change))
        f=next(x for x in p['fields'] if x['id']=='E17')
        self.assertEqual(f['cell'],'E40');self.assertTrue(p['blocking'])

    def test_preview_detects_moved_wage_by_label_not_old_number(self):
        from profit_model.excel_import import preview_excel, confirmed_excel
        def change(w):
            s=w['成本计算'];s['C17']=None;s['E17']=999
            s['C40']='生产人员工资';s['E40']=2;s['F40']='元/双'
        content=modified(change);p=preview_excel(content)
        wage=next(x for x in p['fields'] if x['id']=='E17')
        self.assertEqual(wage['cell'],'E40');self.assertEqual(wage['value'],'2')
        self.assertNotIn('production_profit',p)
        data=confirmed_excel(content,p['mapping'],p['source_hash'])
        self.assertEqual(data['source_analysis']['metrics']['production_profit'],'1562.50')
        self.assertEqual(data['source_analysis']['import_review']['fields']['E17']['cell'],'E40')

    def test_total_cost_converted_using_order_quantity(self):
        from profit_model.excel_import import preview_excel, confirmed_excel
        content=modified(lambda w:w['成本计算'].__setitem__('G40',200))
        p=preview_excel(content);p['mapping']['fields']['E17']={'cell':'G40','unit':'total'}
        data=confirmed_excel(content,p['mapping'],p['source_hash'])
        self.assertEqual(data['source_analysis']['metrics']['production_profit'],'1562.50')

    def test_total_conversion_preserves_decimal_precision(self):
        from profit_model.excel_import import preview_excel, confirmed_excel
        def change(w):
            w['原订单用量']['H3']=3;w['原订单用量']['H7']=3
            w['生产周期']['F3']=3;w['生产周期']['B8']=11
            w['成本计算']['E8']=3;w['成本计算']['G40']=100
        content=modified(change);p=preview_excel(content)
        p['mapping']['fields']['E17']={'cell':'G40','unit':'total'}
        a=confirmed_excel(content,p['mapping'],p['source_hash'])['source_analysis']
        self.assertEqual(a['metrics']['processing_total'],'100.00')

    def test_ambiguous_label_requires_explicit_mapping(self):
        from profit_model.excel_import import preview_excel, confirmed_excel
        def change(w):
            s=w['成本计算'];s['C17']='生产人员工资';s['C40']='直接人工';s['E40']=2
        content=modified(change);p=preview_excel(content)
        self.assertTrue(p['blocking'])
        with self.assertRaises(ValueError):confirmed_excel(content,{},p['source_hash'])
        p['mapping']['fields']['E17']={'cell':'E40','unit':'per_pair'}
        self.assertEqual(confirmed_excel(content,p['mapping'],p['source_hash'])['source_analysis']['metrics']['production_profit'],'1562.50')

    def test_profit_answer_cannot_be_mapped_as_cost(self):
        from profit_model.excel_import import preview_excel, confirmed_excel
        content=source_bytes();p=preview_excel(content)
        p['mapping']['fields']['E17']={'cell':'H24','unit':'per_pair'}
        with self.assertRaisesRegex(ValueError,'结果|汇总'):confirmed_excel(content,p['mapping'],p['source_hash'])

    def test_moved_subtotal_cannot_be_used_directly_or_indirectly(self):
        from profit_model.excel_import import preview_excel, confirmed_excel
        def change(w):
            s=w['成本计算'];s['E13']=2;s['C40']='加工成本合计';s['E40']='=SUM(E13:E23)';s['G42']='=E40'
        content=modified(change)
        for cell in ('E40','G42'):
            p=preview_excel(content);p['mapping']['fields']['E17']={'cell':cell,'unit':'per_pair'}
            with self.subTest(cell=cell),self.assertRaisesRegex(ValueError,'结果|汇总'):confirmed_excel(content,p['mapping'],p['source_hash'])

    def test_annotated_profit_result_cannot_be_mapped_as_input(self):
        from profit_model.excel_import import preview_excel, confirmed_excel
        def change(w):
            w['成本计算']['C40']='模拟生产毛利（旧表费用口径）';w['成本计算']['E40']=123
        content=modified(change);p=preview_excel(content)
        p['mapping']['fields']['E17']={'cell':'E40','unit':'per_pair'}
        with self.assertRaisesRegex(ValueError,'结果|汇总'):confirmed_excel(content,p['mapping'],p['source_hash'])

    def test_hash_mismatch_and_invalid_cells_rejected(self):
        from profit_model.excel_import import preview_excel, confirmed_excel
        content=source_bytes();p=preview_excel(content)
        with self.assertRaisesRegex(ValueError,'文件'):confirmed_excel(content,p['mapping'],'wrong-hash')
        p['mapping']['fields']['E17']={'cell':'../../secret','unit':'per_pair'}
        with self.assertRaises(ValueError):confirmed_excel(content,p['mapping'],p['source_hash'])

    def test_reused_mapping_reads_new_values_and_preserves_missing_fee(self):
        from profit_model.excel_import import preview_excel, confirmed_excel
        first=modified(lambda w:w['成本计算'].__setitem__('G40',2))
        p=preview_excel(first);p['mapping']['fields']['E17']={'cell':'G40','unit':'per_pair'}
        new=modified(lambda w:w['成本计算'].__setitem__('G40',3))
        q=preview_excel(new,p['mapping'])
        self.assertEqual(next(x['value'] for x in q['fields'] if x['id']=='E17'),'3')
        self.assertEqual(confirmed_excel(new,p['mapping'],q['source_hash'])['source_analysis']['metrics']['production_profit'],'1462.50')
        missing=modified(lambda w:w['成本计算'].__setitem__('E17',None))
        r=preview_excel(missing)
        self.assertIsNone(confirmed_excel(missing,r['mapping'],r['source_hash'])['source_analysis']['metrics']['production_profit'])


class PreviewApiTests(unittest.TestCase):
    def setUp(self):
        from server import create_server
        self.tmp=tempfile.TemporaryDirectory();self.storage=Path(self.tmp.name)
        self.server=create_server(0,self.storage,Mock());self.thread=threading.Thread(target=self.server.serve_forever,daemon=True);self.thread.start()
        self.url=f'http://127.0.0.1:{self.server.server_port}'
        self.opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))

    def tearDown(self):
        self.server.shutdown();self.server.server_close();self.thread.join();self.tmp.cleanup()

    def post(self,path,payload):
        req=urllib.request.Request(self.url+path,data=json.dumps(payload).encode(),headers={'Content-Type':'application/json'})
        with self.opener.open(req) as r:return json.load(r)

    def test_preview_is_read_only_and_confirmation_required(self):
        payload={'workbook_base64':base64.b64encode(source_bytes()).decode()}
        p=self.post('/api/preview-xlsx',payload)
        self.assertEqual(list(self.storage.iterdir()),[])
        with self.assertRaises(urllib.error.HTTPError):self.post('/api/import-reviewed-xlsx',dict(payload,mapping=p['mapping'],source_hash=p['source_hash']))
        r=self.post('/api/import-reviewed-xlsx',dict(payload,mapping=p['mapping'],source_hash=p['source_hash'],confirmed=True))
        self.assertEqual(r['source_analysis']['metrics']['production_profit'],'1662.50')
        self.assertTrue((self.storage/(r['snapshot']['snapshot_id']+'.json')).is_file())

    def test_automatic_import_saves_only_ready_result(self):
        from test_auto_import import labeled_source
        payload={'workbook_base64':base64.b64encode(source_bytes()).decode()}
        held=self.post('/api/auto-import-xlsx',payload)
        self.assertEqual(held['status'],'needs_review');self.assertEqual(list(self.storage.iterdir()),[])
        payload['workbook_base64']=base64.b64encode(labeled_source()).decode()
        done=self.post('/api/auto-import-xlsx',payload)
        self.assertEqual(done['status'],'imported')
        self.assertEqual(done['result']['snapshot']['import_review']['mode'],'automatic')
        self.assertEqual(len(list(self.storage.glob('*.json'))),2)
