"""Inspectable Excel import: label proposals, explicit mappings and original-file provenance."""
import hashlib
import io
import re
import zipfile
from datetime import date,datetime
from decimal import Decimal, localcontext
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple
from .eva_source import FormulaReader, SHEETS, read_eva_workbook
from .files import read_workbook, HEADERS
from .trust import review_source

FIELDS=[('E4','订单单价','订单',('订单单价','合同单价','报价','单价')),
        ('E5','每双折扣','订单',('订单折扣','每双折扣','折扣金额','折扣')),
        ('E8','订单数量','订单',('订单数量','订单总数量','接单数量'))]
for row,label,aliases in [
    (13,'电费',('电费','电力费用')),(14,'天然气',('天然气','天然气费','燃气费')),
    (15,'水费',('水费',)),(16,'场地',('场地','场地费','租金')),
    (17,'生产人员工资',('生产人员工资','生产工资','直接人工','直接人工成本')),
    (18,'生产管理工资',('管理人员工资','生产管理工资')),
    (19,'前期开发成本',('前期开发成本','开发费','开发成本')),
    (20,'机器折旧',('机器折旧','设备折旧')),(21,'生产物资领用',('物资领用','生产物资领用')),
    (22,'生产其他费用',('其他费用','生产其他费用')),(23,'包装费',('包装费','包装费用')),
    (27,'行政人员',('行政人员','行政工资','行政人员工资')),
    (28,'固定资产摊分',('固定资产摊分','固定资产分摊')),
    (29,'周边管理费',('周边管理费','行政管理费')),
    (30,'行政物资领用',('物资领用','行政物资领用')),
    (31,'行政其他费用',('其他费用','行政其他费用')),(32,'运费',('运费','运输费'))]:
    FIELDS.append(('E'+str(row),label,'加工成本' if row<24 else '行政成本',aliases))
FIELD_IDS={x[0] for x in FIELDS}
PROTECTED=('H5','H10','E9','E10','H23','H24','H32','H33')


def norm(value):return re.sub(r'[\s\-—_：:（）()]+','',str(value or '')).lower()


def input_value(value):
    if value is None or isinstance(value,bool):return False
    if isinstance(value,str) and value.startswith('='):return True
    try:return Decimal(str(value)).is_finite()
    except Exception:return False


def books(content):
    if not isinstance(content,bytes) or not 0<len(content)<=10*1024*1024:raise ValueError('Excel文件须在0至10MB之间')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(x.file_size for x in z.infolist())>50*1024*1024:raise ValueError('Excel解压大小超过50MB')
        w=load_workbook(io.BytesIO(content),data_only=False,keep_links=False)
        c=load_workbook(io.BytesIO(content),data_only=True,keep_links=False)
        if len(w.sheetnames)>50 or any(s.max_row*s.max_column>200000 for s in w):
            w.close();c.close();raise ValueError('工作表范围过大，请删除多余空白行列后导入')
        return w,c
    except (OSError,zipfile.BadZipFile,KeyError):raise ValueError('无法读取Excel文件') from None


def locate_sheet(w,mapping):
    selected=mapping.get('cost_sheet')
    if selected:
        if not isinstance(selected,str) or selected not in w.sheetnames:raise ValueError('请选择有效的成本工作表')
        return selected
    if '成本计算' in w.sheetnames:return '成本计算'
    matches=[s.title for s in w if '成本' in s.title or str(s['A1'].value or '').startswith('订单成本')]
    return matches[0] if len(matches)==1 else None


def layout_info(w,cost_sheet=None):
    """Bind reused mappings to labels, units, positions and cell types, never amounts."""
    structure=[]
    for s in w:
        if s.title not in SHEETS[:3] and s.title!=cost_sheet and not ('metadata' in w.sheetnames):continue
        for row in s:
            for cell in row:
                v=cell.value
                if v is None:continue
                if isinstance(v,(date,datetime)):kind='date'
                elif input_value(v):kind='formula' if str(v).startswith('=') else 'number'
                else:
                    header=(s.title in SHEETS[:2] and (cell.row==5 or cell.column==6 and 5<=cell.row<=15)) or s.title=='生产周期' and cell.row in (5,6)
                    kind='text:'+norm(v) if header or s.title==cost_sheet and cell.coordinate!='A1' else 'text'
                structure.append((s.title,cell.coordinate,kind))
    signature=hashlib.sha256(repr(structure).encode('utf-8')).hexdigest()
    reasons=[];overflow=[]
    if not all(x in w.sheetnames for x in SHEETS[:3]):return signature,['材料或日报布局尚未确认'],overflow
    for name in SHEETS[:2]:
        s=w[name]
        expected={'B5':{'材料编号','材料名称'},'C5':{'材料成本','材料单价'},'D5':{'配方用量'}}
        if any(norm(s[cell].value) not in values for cell,values in expected.items()):reasons.append(name+'的配方表头位置不能自动确认')
        for row in (5,9,13):
            if norm(s[f'F{row}'].value)!='码号' or 'sop' not in norm(s[f'F{row+1}'].value) or norm(s[f'F{row+2}'].value)!='订单数量':
                reasons.append(name+'的SOP和分码数量区域不能自动确认');break
        for row in range(18,s.max_row+1):
            if s[f'B{row}'].value is not None and input_value(s[f'C{row}'].value) and input_value(s[f'D{row}'].value):
                overflow.append(name+'在第'+str(row)+'行存在模板范围外的材料明细，需调整布局后再导入');break
    s=w['生产周期']
    if any(norm(s[cell].value) not in allowed for cell,allowed in {'A5':{'生产日期','日期'},'B5':{'生产数量','生产数'},'F5':{'不良数','不良数量'},'J5':{'生产欠数','剩余数量'}}.items()):reasons.append('生产日报的日期、生产、不良及剩余数量表头不能自动确认')
    for a,b,c in [('B','F','J'),('C','G','K'),('D','H','L')]:
        if not s[a+'6'].value or not s[a+'6'].value==s[b+'6'].value==s[c+'6'].value:reasons.append('生产日报的分码列对应关系不能自动确认');break
    for row in range(44,s.max_row+1):
        value=s[f'A{row}'].value
        if isinstance(value,(date,datetime)) or input_value(value) or isinstance(value,str) and re.match(r'^\s*\d{4}[-/年.]\d{1,2}',value):
            overflow.append('生产日报第'+str(row)+'行超出当前读取区域，需扩展适配后再导入');break
    return signature,list(dict.fromkeys(reasons)),overflow


def inspect_eva(w,c,mapping):
    if not isinstance(mapping,dict) or not isinstance(mapping.get('fields',{}),dict):raise ValueError('导入方案格式无效')
    sheet=locate_sheet(w,mapping)
    if not sheet:return {'kind':'eva','fields':[],'mapping':{'fields':{}},'blocking':['无法唯一识别成本工作表，请选择后重新预览'],'sheet_names':w.sheetnames,'regions':[]}
    supplied=mapping.get('fields',{})
    if set(supplied)-FIELD_IDS:raise ValueError('导入方案含未知字段')
    s=w[sheet];reader=FormulaReader(w,c)
    reader.excluded={(sheet,x) for x in PROTECTED}
    reader.excluded.update((name,x) for name in SHEETS[:2] for x in ('B19','B20','K3','D18','E18'))
    # Profit rows may move as well. Their numeric/result cells must stay out of inputs.
    for row in s:
        for cell in row:
            text=norm(cell.value)
            subtotal=any(x in text for x in ('合计','汇总','总额','总计')) and any(x in text for x in ('加工','行政','材料','生产成本','总成本','费用','成本'))
            profit=not str(cell.value or '').startswith('=') and any(x in text for x in ('毛利','净利','纯利'))
            if profit or text in ('加工成本','行政成本','材料成本','生产成本','总成本','总费用','合计','汇总') or subtotal:
                reader.excluded.update((sheet,s.cell(cell.row,j).coordinate) for j in range(cell.column+1,min(cell.column+6,s.max_column+1)))
    labels=[];group='订单'
    for row in s:
        for cell in row:
            text=norm(cell.value)
            if text in ('加工成本','生产成本','加工费用'):group='加工成本'
            if text in ('行政成本','行政费用','经营费用'):group='行政成本'
            if text:labels.append((cell,group,text))
    fields=[];blocking=[];used={};resolved={}
    for fid,label,section,aliases in FIELDS:
        candidates=[]
        for cell,g,text in labels:
            if text not in {norm(x) for x in aliases}:continue
            if label in ('生产物资领用','生产其他费用','行政物资领用','行政其他费用') and g!=section:continue
            values=[s.cell(cell.row,j) for j in range(cell.column+1,min(cell.column+4,164)) if input_value(s.cell(cell.row,j).value)]
            candidates.extend(x.coordinate for x in values)
        candidates=list(dict.fromkeys(candidates))
        explicit=fid in supplied
        entry=supplied.get(fid,{})
        if not isinstance(entry,dict):raise ValueError('字段映射格式无效')
        cell=entry.get('cell',candidates[0] if len(candidates)==1 else '' if len(candidates)>1 else fid)
        status='manual' if explicit else 'label' if len(candidates)==1 else 'ambiguous' if candidates else 'template'
        unit=entry.get('unit','quantity' if fid=='E8' else 'per_pair')
        allowed={'quantity'} if fid=='E8' else {'per_pair'} if fid=='E4' else {'per_pair','total'}
        raw=None;value=None;error=None;source_unit='';source_label=''
        if not isinstance(cell,str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,5}',cell):
            error='存在多个候选位置，请明确选择' if not cell else '单元格位置无效'
        else:
            r,col=coordinate_to_tuple(cell)
            if r>10000 or col>163: error='位置超出本次支持范围'
            elif cell in used:error='与'+used[cell]+'使用同一数值位置，请核对'
            elif (sheet,cell) in reader.excluded:error='结果或汇总单元格不能作为计算输入'
            else:
                used[cell]=label;raw=s[cell].value
                left=[x.value for x in s[cell].parent[s[cell].row][:col-1] if isinstance(x.value,str) and not input_value(x.value)]
                source_label=str(left[-1]) if left else ''
                source_unit=str(s.cell(r,col+1).value or '')
                if not explicit and norm(source_unit) in ('元/单','元/整单','元/订单','元每单'):unit='total'
                elif not explicit and source_unit and ('元' in source_unit or '/' in source_unit or '%' in source_unit) and norm(source_unit) not in ('/双','元/双','元每双'):
                    error='原单位不能自动换算，请明确确认单位，不能直接将小时或百分比当作每双费用'
                if unit not in allowed:error='字段单位无效'
                else:
                    try:value=None if raw is None else reader.value(sheet,cell)
                    except (ValueError,ArithmeticError) as e:error=str(e)
        if fid in ('E4','E5','E8') and value is None and not error:error='订单数量、单价和折扣必须明确提供，零折扣请填0'
        if value is not None and value<0 and fid!='E5':error='数量或费用不能为负数'
        if fid=='E8' and value is not None and value<=0:error='订单数量必须大于零'
        if error:blocking.append(label+'：'+error)
        resolved[fid]={'cell':cell,'unit':unit}
        fields.append({'id':fid,'label':label,'group':section,'sheet':sheet,'cell':cell,'unit':unit,'source_unit':source_unit,'source_label':source_label,'raw':str(raw) if raw is not None else None,'value':None if value is None else str(value),'status':status,'candidates':candidates,'error':error})
    regions=[]
    for name,area in [('原订单用量','B3:K17'),('补单用量','B3:K17'),('生产周期','A6:L43')]:
        regions.append({'sheet':name,'range':area,'rows':[[{'cell':cell.coordinate,'value':str(cell.value) if cell.value is not None else ''} for cell in row] for row in w[name][area]]})
    result={'kind':'eva','fields':fields,'mapping':{'cost_sheet':sheet,'fields':resolved},'blocking':blocking,'sheet_names':w.sheetnames,'regions':regions,
            'warnings':['费用项优先按标签识别；标为模板位置的字段仍需人工核对。','材料配方、SOP分码和生产日报目前仍按展示的模板区域读取；若这些区域发生结构变化，请先调整表格或另行适配。','没有单价的费用保留未知；整单费用按本单数量换算为每双费用。','公式按原文件的引用计算；调整字段位置映射不会自动修改原表公式。']}
    return result,reader


def preview_excel(content,mapping=None):
    mapping={} if mapping is None else mapping
    if not isinstance(mapping,dict):raise ValueError('导入方案格式无效')
    w,c=books(content)
    try:
        if 'metadata' in w.sheetnames and 'orders' in w.sheetnames:
            data=read_workbook(content)
            signature,_,_=layout_info(w)
            reasons=[]
            for name,headers in HEADERS.items():
                if name in w.sheetnames and {str(x.value or '') for x in w[name][1] if x.value is not None}!=set(headers):reasons.append(name+'的表头与标准模板不同')
            if not data['orders']:reasons.append('没有可导入的订单')
            return {'kind':'standard','source_hash':hashlib.sha256(content).hexdigest(),'mapping':{},'blocking':[],
                    'layout_signature':signature,'auto_eligible':not reasons,'review_reasons':reasons,
                    'fields':[],'sheet_names':w.sheetnames,'tables':[{'name':s.title,'rows':max(0,s.max_row-1),'headers':[str(x.value or '') for x in s[1]],'samples':[[str(x.value) if x.value is not None else '' for x in row] for row in s.iter_rows(min_row=2,max_row=min(s.max_row,4))]} for s in w],
                    'warnings':['标准模板按表头读取，允许列顺序变化。确认后将进行数据完整性检查和计算。'],'regions':[]}
        if not all(name in w.sheetnames for name in SHEETS[:3]):
            raise ValueError('尚不支持此表格结构。请使用标准模板，或保留原订单用量、补单用量、生产周期三张明细表；任意材料/日报布局仍需适配。')
        inspected=inspect_eva(w,c,mapping)
        result=inspected[0] if isinstance(inspected,tuple) else inspected
        result['source_hash']=hashlib.sha256(content).hexdigest()
        signature,reasons,overflow=layout_info(w,result['mapping'].get('cost_sheet'))
        result['layout_signature']=signature;result['blocking'].extend(overflow)
        for field in result['fields']:
            uncertain_position=field['status'] in ('template','ambiguous') or field['status']=='manual' and (len(field['candidates'])!=1 or field['cell']!=field['candidates'][0])
            uncertain_unit=(norm(field['source_unit']) not in ('','双')) if field['id']=='E8' else norm(field['source_unit']) not in ('/双','元/双','元每双','元/单','元/整单','元/订单','元每单')
            field['requires_review']=bool(field['error'] or field['value'] is None or uncertain_position or uncertain_unit)
            if uncertain_position:reasons.append(field['label']+'的读取位置需要确认')
            if field['value'] is None:reasons.append(field['label']+'缺少可用数值')
            if field['id']!='E8' and norm(field['source_unit']) not in ('/双','元/双','元每双','元/单','元/整单','元/订单','元每单'):reasons.append(field['label']+'的单位需要确认')
            if field['id']=='E8' and norm(field['source_unit']) not in ('','双'):reasons.append('订单数量单位需要确认')
        result['auto_eligible']=not reasons and not result['blocking']
        result['review_reasons']=list(dict.fromkeys(reasons+result['blocking']))
        result['overflow_reasons']=overflow
        return result
    finally:w.close();c.close()


def auto_excel(content,profiles=None):
    """Auto-adapt only a known layout or an explicitly approved unchanged profile."""
    p=preview_excel(content);mode='automatic';profile_name=None
    if not p['auto_eligible']:
        if profiles is not None and (not isinstance(profiles,list) or len(profiles)>20):raise ValueError('导入方案列表无效')
        for profile in profiles or []:
            if not isinstance(profile,dict) or profile.get('approved') is not True or profile.get('layout_signature')!=p['layout_signature']:continue
            candidate=preview_excel(content,profile.get('mapping'))
            if candidate['blocking'] or candidate.get('overflow_reasons') or any(f['value'] is None for f in candidate['fields']):continue
            p=candidate;mode='automatic_profile';profile_name=str(profile.get('name','已确认方案'))[:60];break
        else:
            return {'status':'needs_review','preview':p,'reasons':p['review_reasons']}
    try:data=confirmed_excel(content,p['mapping'],p['source_hash'])
    except ValueError as e:
        return {'status':'needs_review','preview':p,'reasons':[str(e)]}
    review=data.get('source_analysis',{}).get('import_review') or {'source_hash':p['source_hash'],'fields':{}}
    review.update(mode=mode,confirmed=False,profile_name=profile_name,layout_signature=p['layout_signature'],scope='程序已按字段、单位和表格结构自动适配；未据此核实业务真实性。')
    if data.get('source_analysis'):
        data['source_analysis']['import_review']=review
        for record in data['costs']:
            record['source']=record.get('source','').replace('位置和单位已确认','自动适配位置和单位')
    data['metadata']['import_review']=review
    return {'status':'ready','data':data,'mode':mode,'profile_name':profile_name}


def confirmed_excel(content,mapping,source_hash):
    if hashlib.sha256(content).hexdigest()!=source_hash:raise ValueError('文件与预览版本不一致，请重新预览文件')
    p=preview_excel(content,mapping)
    if p['blocking']:raise ValueError('；'.join(p['blocking']))
    if p['kind']=='standard':return read_workbook(content)
    w,c=books(content)
    try:
        _,reader=inspect_eva(w,c,mapping)
        actual=p['mapping']['cost_sheet'];s=w[actual]
        qty=Decimal(next(x['value'] for x in p['fields'] if x['id']=='E8'))
        values={}
        with localcontext() as ctx:
            ctx.prec=80
            for f in p['fields']:
                value=None if f['value'] is None else Decimal(f['value'])
                values[f['id']]=None if value is None else value/qty if f['unit']=='total' else value
        external={};traces={}
        for f in p['fields']:
            key=(actual,f['cell'])
            if key in reader.external:
                item=dict(reader.external[key],sheet='成本计算',cell=f['id'],source_sheet=actual,source_cell=f['cell'])
                external['成本计算',f['id']]=item
            if key in reader.traces:traces[actual,f['cell']]=reader.traces[key]
        moved=actual!='成本计算' or any(f['cell']!=f['id'] for f in p['fields'])
        if actual!='成本计算':
            if '成本计算' in w.sheetnames:raise ValueError('成本工作表名称冲突，请核对选择')
            s.title='成本计算';c[actual].title='成本计算'
        for f in p['fields']:
            s[f['id']]=None if values[f['id']] is None else str(values[f['id']])
            if f['id'] not in ('E4','E5','E8'):
                row=coordinate_to_tuple(f['id'])[0];s[f'C{row}']=f['label'];s[f'F{row}']='元/双'
        if moved:
            # Saved answers at the old positions must not become false comparison values.
            c['成本计算']['H24']=None;c['成本计算']['H33']=None
        data=read_eva_workbook(content,prepared=(w,c,external,traces,{('成本计算',key):value for key,value in values.items()}))
        a=data['source_analysis']
        review={'source_hash':p['source_hash'],'cost_sheet':actual,'fields':{f['id']:dict(p['mapping']['fields'][f['id']],label=f['label'],read_value=f['value'],source_unit=f['source_unit']) for f in p['fields']},'confirmed':True,'mode':'reviewed','layout_signature':p['layout_signature'],'scope':'费用位置和单位已由导入方确认；材料与日报仍使用预览区域，业务真实性未据此核实。'}
        a['import_review']=review
        for row in a['cost_rows']:
            row['source_sheet']=actual;row['source_cell']=p['mapping']['fields'][row['cell']]['cell']
            record=next(x for x in data['costs'] if x['record_id']=='source-'+row['cell'])
            record['source']=f"{actual}!{row['source_cell']} · {row['label']}（位置和单位已确认，业务依据待核验）"
        a['audit']['excluded_cells']=sorted(set(a['audit']['excluded_cells'])|{f'{name}!{cell}' for name,cell in reader.excluded})
        if '模拟数据总表' in w.sheetnames:
            data['metadata']['is_demo']=True
            a['warnings'].insert(0,{'code':'SIMULATION','message':'文件包含模拟数据总表，本次按模拟数据展示，不认定为企业实际利润。'})
        a['trust_review']=review_source(a)
        data['metadata']['import_review']=review
        return data
    finally:w.close();c.close()
