"""Independent theoretical production-margin inputs; never reads actual results."""

import hashlib
import io
import re
import zipfile
from xml.etree import ElementTree
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ROWS = 1000
MAX_COLUMNS = 100
LIMIT = Decimal('1e15')
QUANTUM = Decimal('0.000000000001')
LABELS = {
    'method': '计算方式',
    'business_order_id': '业务订单号',
    'batch_id': '批次',
    'quantity': '订单数量',
    'unit': '计量单位',
    'unit_price': '报价单价',
    'discount_per_pair': '每双折扣',
    'material_cost': '理论材料成本',
    'labor_cost': '理论人工成本',
    'manufacturing_cost': '理论加工制造成本',
    'other_production_cost': '其他生产成本',
    'target_margin_pct': '目标生产毛利率',
    'tax_basis': '税口径',
    'source_note': '来源说明',
    'target_defect_pct': '目标不良率',
    'material_pct': '材料成本比例',
    'production_management_pct': '生产经管比例',
    'production_loss_pct': '生产损耗比例',
    'defect_loss_pct': '不良损耗比例',
    'energy_pct': '能源成本比例',
    'preparation_labor_per_pair': '准备人工单价',
    'press_labor_per_pair': '压制人工单价',
    'company_management_pct': '公司经管比例',
    'administration_pct': '行政费用比例',
}
ALIASES = {key: key for key in LABELS}
ALIASES.update({label: key for key, label in LABELS.items()})
ALIASES.update(
    {
        '方法': 'method',
        '批次号': 'batch_id',
        '数量': 'quantity',
        '单位': 'unit',
        '单价': 'unit_price',
        '材料成本': 'material_cost',
        '人工成本': 'labor_cost',
        '加工/制造成本': 'manufacturing_cost',
        '加工制造成本': 'manufacturing_cost',
        '折扣': 'discount_per_pair',
        '理论生产毛利率': 'target_margin_pct',
        '目标毛利率': 'target_margin_pct',
        '理论毛利率': 'target_margin_pct',
        '材料比例': 'material_pct',
        '生产经管': 'production_management_pct',
        '生产损耗': 'production_loss_pct',
        '不良品损耗': 'defect_loss_pct',
        '不良损耗': 'defect_loss_pct',
        '能源比例': 'energy_pct',
        '准备工价': 'preparation_labor_per_pair',
        '压制工价': 'press_labor_per_pair',
        '公司经管': 'company_management_pct',
        '行政费用': 'administration_pct',
    }
)
COST_FIELDS = (
    'material_cost',
    'labor_cost',
    'manufacturing_cost',
    'other_production_cost',
)
QUOTE_FIELDS = (
    'material_pct',
    'production_management_pct',
    'production_loss_pct',
    'defect_loss_pct',
    'energy_pct',
    'preparation_labor_per_pair',
    'press_labor_per_pair',
    'company_management_pct',
    'administration_pct',
)
NUMERIC_FIELDS = (
    ('quantity', 'unit_price', 'discount_per_pair')
    + COST_FIELDS
    + ('target_margin_pct', 'target_defect_pct')
    + QUOTE_FIELDS
)


def _empty(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _number(value, field, required=False, percentage=False):
    if _empty(value):
        if required:
            raise ValueError(f'{LABELS[field]}不能为空')
        return None
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise ValueError(f'{LABELS[field]}必须是数字')
    text = str(value).strip()
    if percentage and text.endswith('%'):
        text = text[:-1].strip()
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f'{LABELS[field]}必须是数字') from exc
    if not number.is_finite() or number.copy_abs() > LIMIT:
        raise ValueError(f'{LABELS[field]}必须是有限数值，绝对值不超过1e15')
    # Trailing zeroes do not add meaningful input precision.
    with localcontext() as context:
        context.prec = 80
        if number.normalize().as_tuple().exponent < -12:
            raise ValueError(f'{LABELS[field]}最多支持12位小数')
    return number


def _text(payload, field, maximum, default='', required=False):
    value = payload.get(field, default)
    if value is None and not required:
        value = default
    if not isinstance(value, str):
        raise ValueError(f'{LABELS[field]}必须是文本')
    value = value.strip()
    if (required and not value) or len(value) > maximum:
        raise ValueError(
            f'{LABELS[field]}不能为空且不超过{maximum}字'
            if required
            else f'{LABELS[field]}不超过{maximum}字'
        )
    return value


def _format(number):
    if number is None:
        return None
    number = number.quantize(QUANTUM, rounding=ROUND_HALF_UP)
    return '0' if number == 0 else format(number, 'f').rstrip('0').rstrip('.')


def normalize_baseline(payload: dict) -> dict:
    """Validate a form or supported workbook input and return JSON-safe decimals."""
    if not isinstance(payload, dict):
        raise ValueError('理论基线必须是对象')
    with localcontext() as context:
        context.prec = 80
        method = payload.get('method')
        if method not in ('ratio', 'costs', 'quote'):
            raise ValueError('计算方式必须为ratio、costs或quote')
        result = {'method': method}
        for field, maximum, default, required in (
            ('business_order_id', 200, '', True),
            ('batch_id', 100, '', False),
            ('unit', 100, '双', True),
            ('source_note', 1000, '', False),
        ):
            result[field] = _text(payload, field, maximum, default, required)
        tax = payload.get('tax_basis', 'unknown')
        if tax not in ('unknown', 'included', 'excluded', 'ignored'):
            raise ValueError('税口径必须为unknown、included、excluded或ignored')
        result['tax_basis'] = tax
        numbers = {}
        for field in NUMERIC_FIELDS:
            required = (
                (
                    method == 'costs'
                    and field
                    in ('quantity', 'unit_price', 'discount_per_pair') + COST_FIELDS
                )
                or (method == 'ratio' and field == 'target_margin_pct')
                or (
                    method == 'quote'
                    and field in ('quantity', 'unit_price') + QUOTE_FIELDS
                )
            )
            number = _number(
                payload.get(field), field, required, field.endswith('_pct')
            )
            if number is not None:
                if field == 'quantity' and number <= 0:
                    raise ValueError('订单数量必须大于零')
                if method == 'quote' and field == 'unit_price' and number <= 0:
                    raise ValueError('报价单价必须大于零')
                if field not in ('target_margin_pct', 'quantity') and number < 0:
                    raise ValueError(f'{LABELS[field]}不能为负数')
                if field.endswith('_pct') and number > 100:
                    raise ValueError(f'{LABELS[field]}不能超过100%')
            numbers[field] = number
        revenue = cost = profit = operating_expense = operating_profit = None
        if method == 'quote':
            revenue = numbers['quantity'] * numbers['unit_price']
            numbers['material_cost'] = revenue * numbers['material_pct'] / 100
            numbers['labor_cost'] = (
                numbers['quantity']
                * (
                    numbers['preparation_labor_per_pair']
                    + numbers['press_labor_per_pair']
                )
                + revenue * numbers['production_management_pct'] / 100
            )
            numbers['manufacturing_cost'] = revenue * numbers['energy_pct'] / 100
            numbers['other_production_cost'] = (
                revenue
                * (numbers['production_loss_pct'] + numbers['defect_loss_pct'])
                / 100
            )
            cost = sum(numbers[field] for field in COST_FIELDS)
            profit = revenue - cost
            numbers['target_margin_pct'] = profit / revenue * 100
            operating_expense = (
                revenue
                * (numbers['company_management_pct'] + numbers['administration_pct'])
                / 100
            )
            operating_profit = profit - operating_expense
        if method == 'costs':
            revenue = numbers['quantity'] * (
                numbers['unit_price'] - numbers['discount_per_pair']
            )
            if revenue <= 0:
                raise ValueError('理论收入必须大于零，请核对单价和折扣')
            cost = sum(numbers[field] for field in COST_FIELDS)
            profit = revenue - cost
            target = profit / revenue * 100
            if numbers['target_margin_pct'] is not None and abs(
                numbers['target_margin_pct'] - target
            ) > Decimal('0.01'):
                raise ValueError('表内目标与独立计算冲突，请核对理论收入和生产成本')
            numbers['target_margin_pct'] = target
        result.update({key: _format(value) for key, value in numbers.items()})
        result.update(
            revenue=_format(revenue),
            production_cost=_format(cost),
            production_profit=_format(profit),
            operating_expense=_format(operating_expense),
            operating_profit=_format(operating_profit),
        )
        return result


# 表内声明毛利率与按比例推导值的容许差（个百分点）；企业常按一位小数填写。
DECLARED_MARGIN_TOLERANCE = Decimal('0.1')
# 独立推导的计划生产毛利率常见区间；越出时只提示人工核对，不阻止保存。
MARGIN_SANE_MIN = Decimal('0')
MARGIN_SANE_MAX = Decimal('90')


def declared_margin_conflict(payload: dict, baseline: dict):
    """quote 的毛利率由比例独立推导；表内另有声明值时口径可能不同，交人工核对。"""
    if baseline.get('method') != 'quote':
        return None
    declared = payload.get('target_margin_pct')
    derived = baseline.get('target_margin_pct')
    if _empty(declared) or _empty(derived):
        return None
    try:
        declared = _number(declared, 'target_margin_pct', percentage=True)
        difference = abs(declared - Decimal(derived))
    except (ValueError, InvalidOperation):
        return None
    if difference <= DECLARED_MARGIN_TOLERANCE:
        return None
    return (
        f'表内声明的目标毛利率 {_format(declared)}% 与按报价比例推导的'
        f'{_format(Decimal(derived))}% 相差 {_format(difference)} 个百分点，'
        '口径可能不同（例如是否已扣计件人工、比例列是否为百分数），'
        '请人工核对后决定采用哪一个。'
    )


def implausible_margin(baseline: dict):
    """独立推导的毛利率越出常见区间时提示核对；ratio 的声明目标不参与判断。"""
    if baseline.get('method') not in ('quote', 'costs'):
        return None
    value = baseline.get('target_margin_pct')
    if _empty(value):
        return None
    try:
        value = Decimal(value)
    except InvalidOperation:
        return None
    if MARGIN_SANE_MIN < value < MARGIN_SANE_MAX:
        return None
    return (
        f'按表内输入独立推导的计划生产毛利率为 {_format(value)}%，不在常见区间 '
        f'{_format(MARGIN_SANE_MIN)}%～{_format(MARGIN_SANE_MAX)}% 内，'
        '请核对订单数量、报价单价、各成本金额或百分比列的单位与范围后再保存。'
    )


def _field(value):
    if not isinstance(value, str) or value.strip() not in ALIASES:
        raise ValueError(f'不支持的理论字段：{value!s}；请使用公开模板或中英文标准字段')
    return ALIASES[value.strip()]


def _percent_format(cell):
    """Only an unquoted, unescaped percent sign denotes an Excel percentage."""
    fmt = cell.number_format or ''
    quoted = escaped = False
    for char in fmt:
        if escaped:
            escaped = False
        elif char == '\\':
            escaped = True
        elif char == '"':
            quoted = not quoted
        elif char == '%' and not quoted:
            return True
    return False


def _cell_value(cell, field):
    value = cell.value
    if (
        field.endswith('_pct')
        and isinstance(value, (int, float))
        and not isinstance(value, bool)
        and _percent_format(cell)
    ):
        return Decimal(str(value)) * 100
    return value


def preview_theory(content: bytes) -> dict:
    """Read only explicit key/value or one-record header workbooks, without saving."""
    if not isinstance(content, bytes) or not content or len(content) > MAX_UPLOAD_BYTES:
        raise ValueError('理论Excel不能为空且不得超过10MB')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if (
                len(members) > 1000
                or sum(item.file_size for item in members) > MAX_UNCOMPRESSED_BYTES
            ):
                raise ValueError('理论Excel解压后不得超过50MB或包含过多文件')
            # A crafted dimension attribute can hide cells from read_only iteration.
            # Inspect the bounded worksheet XML itself before trusting dimensions.
            for member in members:
                if not (
                    member.filename.startswith('xl/worksheets/')
                    and member.filename.endswith('.xml')
                ):
                    continue
                count = 0
                with archive.open(member) as source:
                    for _, element in ElementTree.iterparse(source, events=('end',)):
                        tag = element.tag.rsplit('}', 1)[-1]
                        if tag == 'f':
                            raise ValueError('理论Excel包含公式，请粘贴为数值后再导入')
                        if tag == 'c':
                            count += 1
                            match = re.fullmatch(
                                r'([A-Z]{1,3})([1-9][0-9]{0,6})', element.get('r', '')
                            )
                            if not match:
                                raise ValueError('理论Excel工作表范围无效')
                            column = 0
                            for letter in match[1]:
                                column = column * 26 + ord(letter) - ord('A') + 1
                            if (
                                int(match[2]) > MAX_ROWS
                                or column > MAX_COLUMNS
                                or count > 10000
                            ):
                                raise ValueError('理论Excel工作表范围过大')
                        element.clear()
        from openpyxl import load_workbook

        book = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('无法读取理论Excel，请上传有效xlsx文件') from exc
    try:
        if len(book.worksheets) > 10:
            raise ValueError('理论Excel最多支持10个工作表')
        for sheet in book.worksheets:
            if (
                sheet.max_row > MAX_ROWS
                or sheet.max_column > MAX_COLUMNS
                or sheet.max_row * sheet.max_column > 10000
            ):
                raise ValueError('理论Excel工作表范围过大')
            for row in sheet.iter_rows():
                if any(cell.data_type == 'f' for cell in row):
                    raise ValueError('理论Excel包含公式，请粘贴为数值后再导入')
        if '理论利润表' in book.sheetnames:
            sheet = book['理论利润表']
        elif len(book.worksheets) == 1:
            sheet = book.worksheets[0]
        else:
            raise ValueError('请使用理论利润表工作表')
        rows = [
            row
            for row in sheet.iter_rows()
            if any(not _empty(cell.value) for cell in row)
        ]
        if not rows:
            raise ValueError('理论利润表不能为空')
        payload = {}

        def add(key, cell):
            field = _field(key)
            if field in payload:
                raise ValueError(f'重复理论字段：{LABELS[field]}')
            payload[field] = _cell_value(cell, field)

        first = [cell.value for cell in rows[0]]
        if first[:2] in (['字段', '数值'], ['field', 'value']):
            if any(not _empty(value) for value in first[2:]):
                raise ValueError('键值模板仅支持字段和数值两列')
            for row in rows[1:]:
                if any(not _empty(cell.value) for cell in row[2:]):
                    raise ValueError('键值模板仅支持字段和数值两列')
                add(row[0].value, row[1])
        else:
            if len(rows) != 2:
                raise ValueError('列式理论表必须首行字段、第二行单条数值')
            for header, cell in zip(rows[0], rows[1]):
                if _empty(header.value) and _empty(cell.value):
                    continue
                add(header.value, cell)
        baseline = normalize_baseline(payload)
        warnings = [
            '仅支持公开模板键值表或首行标准字段、第二行单条数值；请核对业务订单、费用范围和税口径。'
        ]
        if baseline['method'] == 'ratio':
            warnings.append('比例目标未反推收入或成本；缺少依据的金额保持为空。')
        if baseline['tax_basis'] == 'unknown':
            warnings.append('税口径尚未确认。')
        if baseline['method'] == 'quote':
            warnings.append(
                '报价比例核价：生产经管计入人工成本，损耗计入其他生产成本；'
                '公司经管与行政费用单列经营费用，物流不纳入本轮预算。'
            )
        if baseline['tax_basis'] == 'ignored':
            warnings.append('本次不计税，不作税额调整；比较仍需确认费用与税口径范围。')
        return {
            'baseline': baseline,
            'warnings': warnings,
            'source_hash': hashlib.sha256(content).hexdigest(),
        }
    finally:
        book.close()


def theory_template() -> bytes:
    """Return a blank public template: no invented enterprise target or costs."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = '理论利润表'
    sheet.append(['字段', '数值'])
    defaults = {'method': 'ratio', 'unit': '双', 'tax_basis': 'unknown'}
    for field, label in LABELS.items():
        sheet.append([label, defaults.get(field)])
    sheet.column_dimensions['A'].width = 27
    sheet.column_dimensions['B'].width = 40
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='285B50')
    sheet.freeze_panes = 'A2'
    notes = book.create_sheet('填写说明')
    for row in [
        ['填写说明'],
        ['ratio：填写业务订单号和目标生产毛利率；不会反推成本。'],
        [
            'costs：填写业务订单号、数量、报价、每双折扣及四项生产成本，均需明确填写，包括确认为零的费用。'
        ],
        ['成本方法独立计算：（数量×（报价－每双折扣）－四项成本合计）÷收入×100。'],
        [
            'quote：填写数量、报价、全部报价比例及准备/压制人工单价；人工未知不得留空后按零处理，明确为零时填写0。'
        ],
        [
            'quote生产成本：材料比例＋生产经管比例＋生产损耗比例＋不良损耗比例＋能源比例对应收入金额，加两项人工单价×数量。公司经管与行政费用单列经营费用；物流不纳入本轮预算。'
        ],
        [
            '百分数输入示例：20、20%或Excel百分数格式的20%均表示20个百分点；示例不代表企业目标。'
        ],
        [
            '税口径：unknown未确认、included含税、excluded不含税、ignored本次不计税且不作税额调整。'
        ],
        [
            '成本均为同范围总金额；折扣为每计量单位金额。生产成本不含管理、销售、财务费用，勿重复归集。'
        ],
        ['必需业务数据保持空白，请填写后导入；所有工作表均不得含公式。'],
    ]:
        notes.append(row)
    notes.column_dimensions['A'].width = 120
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()
