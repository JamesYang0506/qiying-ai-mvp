"""Read the EVA source layout without overwriting files or silently resolving external links."""
import ast
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP, localcontext
import hashlib
import io
import re
import zipfile
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.datetime import from_excel
from .trust import review_source

SHEETS=('原订单用量','补单用量','生产周期','成本计算')
ZERO=Decimal(0)


def amount(value):
    return None if value is None else str(value.quantize(Decimal('.01'),rounding=ROUND_HALF_UP))


def decimal(value):
    n=Decimal(str(value))
    if not n.is_finite() or n.copy_abs()>Decimal('1e15') or n.as_tuple().exponent < -20:
        raise ValueError('源表数字超出允许范围')
    return n


class FormulaReader:
    """Restricted arithmetic interpreter, not Python eval or an Excel macro runner."""
    def __init__(self,formulas,cached):
        self.formulas=formulas;self.cached=cached;self.values={};self.active=set()
        self.external={};self.traces={}

    def required(self,sheet,cell):
        if self.formulas[sheet][cell].value is None:raise ValueError(f'{sheet}!{cell} 必填数据缺失')
        return self.value(sheet,cell)

    def value(self,sheet,cell):
        cell=cell.replace('$','').upper();key=(sheet,cell)
        if key in getattr(self,'excluded',set()):
            raise ValueError(f'{sheet}!{cell} 是结果/汇总单元格，不能作为独立计算输入')
        if key in self.values:return self.values[key]
        if key in self.active:raise ValueError(f'循环引用 {sheet}!{cell}')
        if sheet not in self.formulas.sheetnames:raise ValueError('引用的源工作表不存在')
        self.active.add(key)
        raw=self.formulas[sheet][cell].value
        try:
            if raw is None:value=ZERO
            elif isinstance(raw,bool):raise ValueError('数字字段不能为布尔值')
            elif isinstance(raw,(int,float,Decimal)):value=decimal(raw)
            elif isinstance(raw,str) and raw.startswith('='):
                cached=self.cached[sheet][cell].value
                if '[' in raw:
                    value=None if cached is None else decimal(cached)
                    self.external[key]={'sheet':sheet,'cell':cell,'formula':raw,'cached_value':None if value is None else str(value),'trust_status':'待核验' if value is not None else '缺失'}
                else:value=self.expression(sheet,raw[1:])
                self.traces[key]={'sheet':sheet,'cell':cell,'formula':raw,
                    'value':None if value is None else str(value),'cached_value':cached,
                    'method':'external_cache' if '[' in raw else 'recalculated'}
            else:
                try:value=decimal(raw)
                except Exception:raise ValueError(f'{sheet}!{cell} 不是数值或受支持公式') from None
            self.values[key]=value
            return value
        finally:self.active.remove(key)

    def expression(self,sheet,expr):
        values={}
        def token(value):
            key='v'+str(len(values));values[key]=value;return key
        def sum_range(match):
            c1,r1,c2,r2=range_boundaries(match.group(1).replace('$',''))
            if (r2-r1+1)*(c2-c1+1)>10000:raise ValueError('源公式范围过大')
            parts=[self.value(sheet,self.formulas[sheet].cell(r,c).coordinate) for r in range(r1,r2+1) for c in range(c1,c2+1)]
            return token(None if any(x is None for x in parts) else sum(parts,ZERO))
        expr=re.sub(r'SUM\(\s*(\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+)\s*\)',sum_range,expr,flags=re.I)
        def ref(match):return token(self.value(match.group(1) or match.group(2) or sheet,match.group(3)))
        expr=re.sub(r"(?:(?:'([^']+)'|([\w\u4e00-\u9fff]+))!)?(\$?[A-Z]{1,3}\$?\d+)",ref,expr)
        def walk(node):
            if isinstance(node,ast.Expression):return walk(node.body)
            if isinstance(node,ast.Name) and node.id in values:return values[node.id]
            if isinstance(node,ast.Constant) and type(node.value) in (int,float):return decimal(node.value)
            if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.UAdd,ast.USub)):
                v=walk(node.operand);return None if v is None else (-v if isinstance(node.op,ast.USub) else v)
            if isinstance(node,ast.BinOp) and isinstance(node.op,(ast.Add,ast.Sub,ast.Mult,ast.Div)):
                a,b=walk(node.left),walk(node.right)
                if a is None or b is None:return None
                if isinstance(node.op,ast.Add):return a+b
                if isinstance(node.op,ast.Sub):return a-b
                if isinstance(node.op,ast.Mult):return a*b
                if b==0:raise ValueError('源公式除数为零')
                return a/b
            raise ValueError('源表公式不在受支持范围，需人工确认后增加适配')
        try:return walk(ast.parse(expr,mode='eval'))
        except SyntaxError:raise ValueError('无法解析源表公式') from None


def read_eva_workbook(content):
    with localcontext() as ctx:
        ctx.prec=80
        data=_read(content)
        a=data['source_analysis'];m=a['metrics']
        def total(*keys):
            values=[m[k] for k in keys]
            return None if any(x is None for x in values) else amount(sum(map(Decimal,values)))
        a['profit_layers']={
            'production_cost':total('material_total','processing_total'),
            'operating_expense':m['administrative_total'],
            'management_expense':amount(sum((Decimal(x['amount']) for x in a['cost_rows'] if x['group']=='行政成本' and x['category']=='management' and x['amount'] is not None),ZERO)) if all(x['amount'] is not None for x in a['cost_rows'] if x['group']=='行政成本' and x['category']=='management') else None,
            'sales_expense':next(x['amount'] for x in a['cost_rows'] if x['cell']=='E32'),
            'cumulative_cost':total('material_total','processing_total','administrative_total'),
            'production_profit':m['production_profit'],'operating_net_profit':m['management_reference_profit'],
            'classification':'暂沿用源表：加工组计入生产成本；行政组拆为管理费用和运费（销售费用）。开发等归属待核验。',
            'status':'参考试算 · 经营净利费用范围待确认'}
        a['trust_review']=review_source(a)
        return data


def _read(content):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist())>50*1024*1024:raise ValueError('源工作簿过大')
        w=load_workbook(io.BytesIO(content),data_only=False,keep_links=False)
        c=load_workbook(io.BytesIO(content),data_only=True,keep_links=False)
    except Exception:raise ValueError('无法读取中底成本源工作簿') from None
    try:
        if not all(s in w.sheetnames for s in SHEETS):raise ValueError('中底成本源表需要四张规定工作表')
        r=FormulaReader(w,c)
        r.excluded={('成本计算',cell) for cell in ('H5','H10','E9','E10','H23','H24','H32','H33')}
        r.excluded.update((sheet,cell) for sheet in SHEETS[:2] for cell in ('B19','B20','K3','D18','E18'))
        qty=r.required('成本计算','E8')
        if qty<=0:raise ValueError('源表订单数量必须大于零')
        product=str(w['原订单用量']['B3'].value or '')
        if not product or any(str(w[s]['B3'].value)!=product for s in SHEETS[:3]):raise ValueError('源表型体关联不一致')
        if any(r.required(s,cell)!=qty for s,cell in [('原订单用量','H3'),('生产周期','F3')]):raise ValueError('源表订单数量不一致')
        oid='SRC-'+product
        warnings=[{'code':'TAX','message':'用户确认税口径待确认。按源表报价复算，不自动除税。'},
            {'code':'IDENTITY','message':f'{product} 是源表订单型体；正式订单主键、客户及订单行标识尚待提供。'},
            {'code':'MATERIAL_BASIS','message':'材料用量由 SOP 重量×数量推算，未证明为实际净耗用；损耗率、领补退料及期末在制品待确认。'},
            {'code':'LABOR','message':'生产工资为源表单位分摊费用，尚无工段计件单价和应计薪数量依据。'},
            {'code':'MFG','message':'制造费用池M、全期工时H及订单工时h未提供，源表缓存费用不等于已验证的工时分摊。'},
            {'code':'PERIOD','message':'费用对应期间及完整分母待确认；生产跨月时须分期核算，不按文件上传日期替代费用期间。'},
            {'code':'NET','message':'行政费用归属和分摊口径待确认，正式经营净利模块关闭；源表“纯利”仅列参考试算。'}]
        for row in list(range(13,24))+list(range(27,33)):
            cell=f'E{row}'
            if w['成本计算'][cell].value is None:
                r.values['成本计算',cell]=None
                warnings.append({'code':'MISSING_RATE','message':f'成本计算!{cell} 单位费用缺失，对应试算利润留空，不填零。'})
        material_info=[];size_quantities={}
        for sheet in SHEETS[:2]:
            recipe=[]
            for row in range(6,18):
                if any(w[sheet][f'{col}{row}'].value is not None for col in ('B','C','D')):
                    unit_price=r.required(sheet,f'C{row}');usage=r.required(sheet,f'D{row}')
                    if unit_price is None or usage is None or unit_price<0 or usage<0:raise ValueError('配方单价或用量缺失/为负')
                    recipe.append({'material':str(w[sheet][f'B{row}'].value or row),'unit_price':str(unit_price),'usage':str(usage),'source':f'{sheet}!C{row}/D{row}'})
            sizes=[]
            for qrow in (7,11,15):
                for col in range(7,12):
                    qcell=w[sheet].cell(qrow,col);weight=w[sheet].cell(qrow-1,col)
                    if qcell.value is None and weight.value is None:continue
                    q=r.required(sheet,qcell.coordinate);kg=r.required(sheet,weight.coordinate)
                    if q<0 or kg<0:raise ValueError('SOP重量和数量不能为负')
                    size=str(w[sheet].cell(qrow-2,col).value or qcell.coordinate)
                    sizes.append({'size':size,'quantity':str(q),'weight_per_pair':str(kg),'weight_total':str(q*kg),'source':f'{sheet}!{weight.coordinate}/{qcell.coordinate}'})
            declared=r.required(sheet,'H3')
            if sum(Decimal(x['quantity']) for x in sizes)!=declared:raise ValueError(f'{sheet} 分码数量与头部数量不一致')
            recipe_weight=sum((Decimal(x['usage']) for x in recipe),ZERO)
            if recipe_weight<=0:raise ValueError('配方总用量必须大于零')
            price=sum((Decimal(x['unit_price'])*Decimal(x['usage']) for x in recipe),ZERO)/recipe_weight
            weight=sum((Decimal(x['weight_total']) for x in sizes),ZERO)
            total=price*weight
            material_info.append({'sheet':sheet,'quantity':str(declared),'weight':str(weight),'price_per_kg':str(price),
                'amount':amount(total),'amount_exact':str(total),'sizes':sizes,'recipe':recipe,'source':f'{sheet} 配方C/D列及分码SOP重量×数量'})
            size_quantities[sheet]={x['size']:Decimal(x['quantity']) for x in sizes}

        days=[];by_size=[];gross=ZERO;defects=ZERO;remaining_checks=[]
        previous={x['size']:Decimal(x['quantity']) for x in material_info[0]['sizes']}
        for row in range(8,44):
            rawdate=w['生产周期'][f'A{row}'].value
            if rawdate is None:continue
            dt=rawdate if isinstance(rawdate,(date,datetime)) else from_excel(rawdate,w.epoch)
            day=dt.date().isoformat() if isinstance(dt,datetime) else dt.isoformat()
            produced_parts=[r.required('生产周期',f'{col}{row}') for col in ('B','C','D')]
            bad_parts=[r.required('生产周期',f'{col}{row}') for col in ('F','G','H')]
            if any(x<0 for x in produced_parts+bad_parts) or any(b>p for b,p in zip(bad_parts,produced_parts)):
                raise ValueError(f'生产周期第{row}行分码数量/不良数异常')
            produced=sum(produced_parts);bad=sum(bad_parts)
            if produced<0 or bad<0 or bad>produced:raise ValueError('生产数量或不良数量范围异常')
            gross+=produced;defects+=bad
            days.append({'date':day,'produced':str(produced),'defects':str(bad),'qualified':str(produced-bad)})
            for idx,(pcol,remain_col) in enumerate(zip(('B','C','D'),('J','K','L'))):
                size=str(w['生产周期'][f'{pcol}6'].value)
                expected=previous.get(size,ZERO)-(produced_parts[idx]-bad_parts[idx])
                ref=f'{remain_col}{row}'
                reported=None if w['生产周期'][ref].value is None else r.value('生产周期',ref)
                remaining_checks.append({'date':day,'size':size,'cell':ref,'expected':str(expected),
                    'reported':None if reported is None else str(reported),'difference':None if reported is None else str(reported-expected)})
                previous[size]=expected
        if not days:raise ValueError('缺少生产日报日期与数量')
        for pcol,dcol in zip(('B','C','D'),('F','G','H')):
            size=str(w['生产周期'][f'{pcol}6'].value)
            produced=sum(r.required('生产周期',f'{pcol}{row}') for row in range(8,44) if w['生产周期'][f'A{row}'].value is not None)
            bad=sum(r.required('生产周期',f'{dcol}{row}') for row in range(8,44) if w['生产周期'][f'A{row}'].value is not None)
            replenish=size_quantities['补单用量'].get(size,ZERO)
            by_size.append({'size':size,'produced':str(produced),'defects':str(bad),'replenishment':str(replenish),'difference':str(bad-replenish)})
        if gross-defects!=qty:warnings.append({'code':'QUALIFIED','message':'日报合格数量与订单数量不一致，不能视为整单已完成。'})
        replenish=r.required('补单用量','H3');difference=defects-replenish
        if difference!=0:warnings.append({'code':'REPLENISHMENT','message':f'生产不良合计{defects}双，补单材料表{replenish}双，相差{difference}双；不自动补增材料。'})

        price=r.required('成本计算','E4');discount=r.required('成本计算','E5');netprice=price-discount
        revenue=netprice*qty if netprice is not None else None
        mat=sum((Decimal(x['amount_exact']) for x in material_info),ZERO)
        processing_parts=[r.value('成本计算',f'E{row}') for row in range(13,24)]
        admin_parts=[r.value('成本计算',f'E{row}') for row in range(27,33)]
        processing=None if any(x is None for x in processing_parts) else sum(processing_parts,ZERO)
        admin=None if any(x is None for x in admin_parts) else sum(admin_parts,ZERO)
        gp=None if processing is None else netprice-mat/qty-processing
        mp=None if gp is None or admin is None else gp-admin
        costs=[{'record_id':'sop-material','order_id':oid,'category':'material','amount':amount(mat),'status':'estimated',
            'stage':'unknown','period':'','method':'source_sop_estimate','source':'配方加权单价 × 分码SOP重量 × 原单/补单数量（模型独立计算）'}]
        rows=[]
        for row in list(range(13,24))+list(range(27,33)):
            ref=f'E{row}';unit=r.value('成本计算',ref)
            category=('labor' if row==17 else 'production_other' if row in (19,21,22) else 'manufacturing') if row<24 else ('sales' if row==32 else 'management')
            external=('成本计算',ref) in r.external
            label=str(w['成本计算'][f'C{row}'].value or ref)
            total=None if unit is None else unit*qty
            entry={'record_id':f'source-{ref}','order_id':oid,'category':category,'amount':amount(total),
                'status':'missing' if total is None else 'estimated','stage':'unknown','period':'',
                'method':'external_cached_rate' if external else 'source_unit_rate','source':f'成本计算!{ref} · {label}（归属待核验）'}
            costs.append(entry)
            rows.append({'label':label,'group':'加工成本' if row<24 else '行政成本','unit_rate':None if unit is None else str(unit),
                'amount':amount(total),'cell':ref,'method':('按表内成本暂算 · 原始依据后续核验' if row in (13,14,15) else '外链缓存待核验' if external else '源表数值/公式复算'),'category':category})
        # Keep traceable per-record cents consistent with the source's aggregate-then-round method.
        for group,total in [('加工成本',None if processing is None else processing*qty),('行政成本',None if admin is None else admin*qty)]:
            if total is None:continue
            rounded=sum(Decimal(x['amount']) for x in rows if x['group']==group and x['amount'] is not None)
            adjustment=Decimal(amount(total))-rounded
            if adjustment:costs.append({'record_id':'round-'+group,'order_id':oid,'category':'production_other' if group=='加工成本' else 'management',
                'amount':amount(adjustment),'status':'estimated','stage':'unknown','period':'','method':'rounding_adjustment','source':group+'汇总后舍入与逐条按分舍入的差额'})
        for entry in r.external.values():
            source_row=next((x for x in rows if x['cell']==entry['cell']),None)
            label=source_row['label'] if source_row else entry['cell']
            message=(f"{label}：已按表内单价计入本次试算（成本计算!{entry['cell']}）。单价原先引用其他工作簿，目前采用文件保存值；其原始费用总额与分摊依据可后续核验。"
                if entry['cached_value'] is not None else f"{label}：表内未保存可用费率，需补金额后才能计算。")
            warnings.append({'code':'EXTERNAL','message':message})
        metrics={'quantity':str(qty),'list_unit_price':str(price),'discount_per_pair':str(discount),'net_unit_price':str(netprice),
            'revenue':amount(revenue),'material_total':amount(mat),'processing_total':amount(None if processing is None else processing*qty),
            'production_profit':amount(None if gp is None else gp*qty),'administrative_total':amount(None if admin is None else admin*qty),
            'management_reference_profit':amount(None if mp is None else mp*qty),
            'unit_production_profit':None if gp is None else str(gp.quantize(Decimal('.000001'),rounding=ROUND_HALF_UP)),
            'unit_management_reference_profit':None if mp is None else str(mp.quantize(Decimal('.000001'),rounding=ROUND_HALF_UP)),
            'production_margin':None if gp is None or netprice is None or netprice<=0 else str(gp/netprice),
            'management_reference_margin':None if mp is None or netprice is None or netprice<=0 else str(mp/netprice)}
        # Only after model metrics are finalized, read saved answers for comparison.
        comparisons=[]
        for label,cell,value in [('毛利润','H24',gp),('扣行政费用后参考利润','H33',mp)]:
            cached=c['成本计算'][cell].value
            try:reference=None if cached is None else decimal(cached)*qty
            except (ValueError,ArithmeticError):reference=None
            calculated=None if value is None else value*qty
            delta=None if reference is None or calculated is None else Decimal(amount(calculated))-Decimal(amount(reference))
            comparisons.append({'label':label,'cell':cell,'calculated':amount(calculated),'reference':amount(reference),'difference':amount(delta)})
        audit={'method':'independent_inputs_v1','excluded_cells':[f'{s}!{cell}' for s,cell in sorted(r.excluded)],
            'comparisons':comparisons,'sequence':'先由明细计算并生成结果，再读取利润保存值作对照；对照值不回写模型。'}
        return {'metadata':{'is_demo':False,'currency':'CNY','cutoff_date':max(x['date'] for x in days),
                'management_enabled':False,'formula_source':'2026-10-13 算法公式与说明汇总','source_hash':hashlib.sha256(content).hexdigest()},
            'orders':[{'order_id':oid,'product':product+' · '+str(w['原订单用量']['E3'].value or ''),'customer':'未提供',
                'quantity':str(qty),'unit':'双','unit_price':str(price),'discount':str((discount*qty).quantize(Decimal('.01'))),
                'mode':'forecast','scope_aligned':False,'revenue_status':'estimated'}],
            'materials':[],'labor':[],'costs':costs,'pools':[],'bases':[],
            'coverage':[{'order_id':oid,'category':cat,'status':'estimated' if cat in ('material','production_other','management','sales') else 'missing'}
                for cat in ('material','labor','manufacturing','production_other','management','sales','finance','order_other')],
            'source_analysis':{'title':'中底成本源表复算（待核验）','order_id':oid,'product':product,'status':'待核验',
                'basis':'由配方、分码SOP重量、数量和逐项费用独立计算；外链费率仍待核验，属于整单参考试算。',
                'audit':audit,
                'unit_notes':[{'cell':'成本计算!E'+str(row),'source_unit':w['成本计算'][f'F{row}'].value,
                    'normalized_unit':'元（材料总额）'} for row in (9,10)],
                'remaining_checks':remaining_checks,
                'source_hash':hashlib.sha256(content).hexdigest(),'metrics':metrics,'materials':material_info,'cost_rows':rows,
                'production':{'start':min(x['date'] for x in days),'end':max(x['date'] for x in days),
                    'gross_quantity':str(gross),'defect_quantity':str(defects),'qualified_quantity':str(gross-defects),
                    'replenishment_quantity':str(replenish),'difference':str(difference),'days':days,'sizes':by_size},
                'external_links':list(r.external.values()),'warnings':warnings,'formula_traces':list(r.traces.values())}}
    finally:w.close();c.close()
