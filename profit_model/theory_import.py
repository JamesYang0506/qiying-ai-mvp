"""Bounded, read-only recognition of theoretical Excel inputs and their provenance."""

import hashlib
import io
import re
import zipfile
from decimal import Decimal
from xml.etree import ElementTree

from .theory import (
    ALIASES,
    COST_FIELDS,
    LABELS,
    MAX_COLUMNS,
    MAX_ROWS,
    MAX_UNCOMPRESSED_BYTES,
    MAX_UPLOAD_BYTES,
    NUMERIC_FIELDS,
    QUOTE_FIELDS,
    _cell_value,
    _empty,
    normalize_baseline,
)


def _label(value):
    if not isinstance(value, str):
        return ''
    return re.sub(r'[\s：:（）()\[\]【】]', '', value).lower()


_ALIASES = {_label(k): v for k, v in ALIASES.items()}
for _key, _values in {
    'business_order_id': (
        '订单号',
        '定单号',
        '订单编号',
        '定单编号',
        '业务定单号',
        '订单',
        '定单',
    ),
    'quantity': ('订单量', '定单数量', '计划数量', '订单数量双', '数量双'),
    'unit_price': (
        '工厂报价',
        '工厂报价元/双',
        '报价元/双',
        '报价单价元/双',
        '单价元/双',
    ),
    'target_margin_pct': (
        '计划毛利率',
        '计划生产毛利率',
        '目标生产毛利率%',
        '目标毛利率%',
        '完整计划毛利率',
        '完整计划生产毛利率',
    ),
    'material_pct': (
        'BOM材料',
        'BOM材料%',
        'BOM材料比例',
        '材料成本%',
        '材料成本比例%',
    ),
    'production_management_pct': ('生产管理工资', '生产管理工资%', '生产管理比例'),
    'production_loss_pct': ('生产损耗%', '生产损耗预算'),
    'defect_loss_pct': ('不良品损耗%', '不良损耗%', '不良品损耗预算'),
    'energy_pct': ('能源', '能源%', '能源成本', '能源成本%', '电汽动力能源'),
    'preparation_labor_per_pair': (
        '备料人工',
        '备料人工元/双',
        '准备人工元/双',
        '备料工价',
        '备料计件人工元/双',
    ),
    'press_labor_per_pair': (
        '油压人工',
        '油压人工元/双',
        '压制人工元/双',
        '油压工价',
        '油压计件人工元/双',
    ),
    'company_management_pct': ('公司管理干部', '公司管理干部%', '公司管理费用比例'),
    'administration_pct': ('行政', '行政%', '行政费用%', '行政人员'),
}.items():
    _ALIASES.update({_label(value): _key for value in _values})


def _field(value):
    key = _label(value)
    field = _ALIASES.get(key)
    if field is None and key.endswith('%'):
        field = _ALIASES.get(key[:-1])
    if field is None:
        total = re.sub(
            r'(总金额|总成本|合计金额|元/双|元／双|元每双|元双|元|双)$', '', key
        )
        if total != key:
            field = _ALIASES.get(total)
    return field


def _load(content):
    """Validate raw worksheet cells before trusting potentially forged dimensions."""
    if not isinstance(content, bytes) or not content or len(content) > MAX_UPLOAD_BYTES:
        raise ValueError('理论Excel不能为空且不得超过10MB')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if (
                len(members) > 1000
                or sum(m.file_size for m in members) > MAX_UNCOMPRESSED_BYTES
            ):
                raise ValueError('理论Excel解压后不得超过50MB或包含过多文件')
            for member in members:
                if not (
                    member.filename.startswith('xl/worksheets/')
                    and member.filename.endswith('.xml')
                ):
                    continue
                count = 0
                with archive.open(member) as source:
                    for _, element in ElementTree.iterparse(source, events=('end',)):
                        tag = element.tag.rsplit('}', 1)[-1]
                        if tag == 'dimension':
                            from openpyxl.utils.cell import range_boundaries

                            try:
                                _, _, max_column, max_row = range_boundaries(
                                    element.get('ref', '')
                                )
                                if not max_column or not max_row:
                                    raise ValueError('理论Excel工作表范围无效')
                                if (
                                    max_row > MAX_ROWS
                                    or max_column > MAX_COLUMNS
                                    or max_row * max_column > 10000
                                ):
                                    raise ValueError('理论Excel工作表范围过大')
                            except (TypeError, ValueError) as exc:
                                raise ValueError(
                                    '理论Excel工作表范围无效或过大'
                                ) from exc
                        if tag == 'c':
                            count += 1
                            match = re.fullmatch(
                                r'([A-Z]{1,3})([1-9][0-9]{0,6})', element.get('r', '')
                            )
                            if not match:
                                raise ValueError('理论Excel工作表范围无效')
                            column = 0
                            for letter in match[1]:
                                column = column * 26 + ord(letter) - ord('A') + 1
                            if (
                                int(match[2]) > MAX_ROWS
                                or column > MAX_COLUMNS
                                or count > 10000
                            ):
                                raise ValueError('理论Excel工作表范围过大')
                        element.clear()
        from openpyxl import load_workbook

        book = load_workbook(
            io.BytesIO(content), read_only=False, data_only=False, keep_links=False
        )
        try:
            if len(book.worksheets) > 10:
                raise ValueError('理论Excel最多支持10个工作表')
            for sheet in book.worksheets:
                if (
                    sheet.max_row > MAX_ROWS
                    or sheet.max_column > MAX_COLUMNS
                    or sheet.max_row * sheet.max_column > 10000
                ):
                    raise ValueError('理论Excel工作表范围过大')
            cached = load_workbook(
                io.BytesIO(content), read_only=False, data_only=True, keep_links=False
            )
            return book, cached
        except Exception:
            book.close()
            raise
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('无法读取理论Excel，请上传有效xlsx文件') from exc


def _layout(sheet):
    rows = list(sheet.iter_rows())
    for row in rows:
        headers = [(cell, _field(cell.value)) for cell in row if _field(cell.value)]
        has_neighbor_value = any(
            any(
                not _empty(c.value) and not _field(c.value)
                for c in row[header.column : next_header.column - 1]
            )
            for (header, _), (next_header, _) in zip(headers, headers[1:])
        )
        if len(headers) >= 2 and not has_neighbor_value:
            records = [
                r
                for r in rows[row[0].row :]
                if any(not _empty(r[c.column - 1].value) for c, _ in headers)
            ]
            return 'records', headers, records
    pairs = []
    for row in rows:
        for cell in row:
            field = _field(cell.value)
            if not field:
                continue
            # The immediate value may intentionally be blank. Never skip it to
            # borrow a derived amount or a note from the following columns.
            if cell.column < len(row):
                value_cell = row[cell.column]
                if field == 'unit' and _empty(value_cell.value):
                    continue
                if not _field(value_cell.value):
                    pairs.append((field, value_cell, str(cell.value)))
    return 'key_value', pairs, []


def _json_value(value):
    if isinstance(value, Decimal):
        text = format(value, 'f')
        return text.rstrip('0').rstrip('.') if '.' in text else text
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _units(cell, label=''):
    """Explicit units are evidence, including Excel display-only formats."""
    units = []
    text = str(cell.value).strip() if cell.value is not None else ''
    match = re.fullmatch(
        r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)\s*(.*)', text.replace(',', '').replace('，', '')
    )
    if match and match[1]:
        units.append(match[1].replace('％', '%').replace('／', '/'))
    fmt = cell.number_format or ''
    if isinstance(_cell_value(cell, 'source_pct'), Decimal):
        units.append('%')
    match = re.search(r'元(?:[/／每]?双)?|双|件|个', fmt)
    if match:
        units.append(match[0].replace('／', '/'))
    match = re.search(r'(%|％|元[/／每]?双|元双|元|双)$', _label(label))
    if match:
        units.append(match[0].replace('％', '%').replace('／', '/'))
    return units


def auto_preview_theory(content: bytes, selection=None) -> dict:
    """Recognize input evidence only; ambiguous fields remain editable and unsaved."""
    selection = {} if selection is None else selection
    if not isinstance(selection, dict) or set(selection) - {
        'sheet',
        'record_row',
        'mapping',
    }:
        raise ValueError('理论表选择必须包含有效工作表、记录行或字段映射')
    mapping = selection.get('mapping', {})
    if not isinstance(mapping, dict) or any(field not in LABELS for field in mapping):
        raise ValueError('理论字段映射无效')
    book, cached = _load(content)
    try:
        layouts = {sheet.title: _layout(sheet) for sheet in book.worksheets}
        counts = {}
        for name, (layout, labels, _) in layouts.items():
            if not labels:
                continue
            # 键值布局的标签是(字段, 值单元格, 标题)，记录布局是(表头单元格, 字段)。
            counts[name] = len(
                {item[0] if layout == 'key_value' else item[1] for item in labels}
            )
        if not counts:
            raise ValueError('未找到可识别的理论字段，请使用公开模板或补充字段标签')
        # 计划表识别的字段最多，对照页通常只剩一两个。字段数领先时直接选中，
        # 并列时仍要求人工选择；选择器照旧列出全部，可覆盖自动结果。
        usable = sorted(counts, key=lambda name: (-counts[name], name))
        warnings = ['自动识别仅供核对；请确认业务订单、成本单位、费用范围和税口径。']
        adaptation = dict(
            sheet=None,
            sheets=[{'name': name} for name in usable],
            records=[],
            fields=[],
            missing=[],
            conflicts=[],
            layout=None,
            notes=[],
        )
        result = dict(
            baseline={},
            warnings=warnings,
            source_hash=hashlib.sha256(content).hexdigest(),
            status='needs_review',
            adaptation=adaptation,
        )
        chosen = selection.get('sheet')
        if chosen is not None and (not isinstance(chosen, str) or chosen not in usable):
            raise ValueError('请选择存在且含理论字段的工作表')
        if chosen is None:
            if len(usable) == 1:
                chosen = usable[0]
            elif counts[usable[0]] > counts[usable[1]]:
                chosen = usable[0]
                adaptation['notes'].append(
                    f'已自动选择识别字段最多的「{chosen}」工作表；如不是计划表，请在上方改选。'
                )
            else:
                adaptation['notes'].append('多个工作表含理论字段，请先选择工作表。')
                return result
        sheet = book[chosen]
        layout, labels, records = layouts[chosen]
        adaptation.update(sheet=chosen, layout=layout)
        pairs = labels
        if layout == 'records':
            order_column = next(
                (cell.column for cell, field in labels if field == 'business_order_id'),
                None,
            )
            adaptation['records'] = [
                {
                    'row': r[0].row,
                    'label': (
                        str(r[order_column - 1].value or '')
                        if order_column
                        else f'第{r[0].row}行'
                    ),
                }
                for r in records
            ]
            record_row = selection.get('record_row')
            if record_row is not None and (
                isinstance(record_row, bool)
                or not isinstance(record_row, int)
                or record_row not in [r[0].row for r in records]
            ):
                raise ValueError('请选择存在的理论记录行')
            if record_row is None:
                if len(records) != 1:
                    adaptation['notes'].append('请选择一条理论记录后继续核对。')
                    return result
                record_row = records[0][0].row
            pairs = [
                (field, sheet.cell(record_row, header.column), str(header.value))
                for header, field in labels
            ]
        elif selection.get('record_row') is not None:
            raise ValueError('当前工作表不是多记录表，无需选择记录行')

        review = False

        def value_for(cell, field, label=''):
            nonlocal review
            if cell.data_type == 'f':
                if '[' in str(cell.value) or re.search(
                    r'(?i)(https?:|file:|\\\\)', str(cell.value)
                ):
                    warnings.append(
                        f'{cell.coordinate}含外部引用公式，未使用其结果，请人工填写。'
                    )
                    review = True
                    return None
                warnings.append(
                    f'{cell.coordinate}为公式，仅读取保存缓存且尚未核实，未执行公式。'
                )
                cell = cached[chosen][cell.coordinate]
            if cell.data_type == 'e':
                warnings.append(f'{cell.coordinate}为Excel错误值，请人工填写。')
                return None
            units = _units(cell, label)
            if field in NUMERIC_FIELDS:
                allowed = (
                    {'%'}
                    if field.endswith('_pct')
                    else (
                        {'双'}
                        if field == 'quantity'
                        else (
                            {'元'}
                            if field in COST_FIELDS
                            else {'元', '元/双', '元每双', '元双'}
                        )
                    )
                )
                if any(unit not in allowed for unit in units):
                    warnings.append(
                        f'{cell.coordinate} {LABELS[field]}的单位与字段不符，请人工核对；未使用此值。'
                    )
                    review = True
                    return None
            value = _cell_value(cell, field)
            if (
                field in ('business_order_id', 'batch_id')
                and isinstance(value, (int, float))
                and not isinstance(value, bool)
            ):
                warnings.append(
                    f'{LABELS[field]}是数字单元格，可能已丢失前导零或精度，请核对原始编号。'
                )
                review = True
                return str(value)
            if field in NUMERIC_FIELDS and isinstance(value, str):
                text = value.strip().replace(',', '').replace('，', '')
                match = re.fullmatch(r'([+-]?(?:\d+(?:\.\d*)?|\.\d+))\s*(.*)', text)
                if match:
                    value = match[1]
                elif not _empty(value):
                    warnings.append(
                        f'{cell.coordinate} {LABELS[field]}不是可用数字，请人工补充；未使用此值。'
                    )
                    review = True
                    return None
            return _json_value(value)

        candidates = {}
        for field, cell, label in pairs:
            # Material percentage has an established quote meaning; other
            # expense percentages cannot safely be assigned a quote category.
            if field == 'material_cost' and '%' in _units(cell, label):
                field = 'material_pct'
            candidate = dict(
                cell=cell.coordinate, value=value_for(cell, field, label), label=label
            )
            candidates.setdefault(field, []).append(candidate)
        for field, coordinate in mapping.items():
            if not isinstance(coordinate, str) or not re.fullmatch(
                r'[A-Z]{1,3}[1-9][0-9]{0,3}', coordinate
            ):
                raise ValueError('字段映射必须是所选工作表内的单元格坐标')
            from openpyxl.utils.cell import coordinate_to_tuple

            row_number, column_number = coordinate_to_tuple(coordinate)
            if row_number > sheet.max_row or column_number > sheet.max_column:
                raise ValueError('字段映射超出所选工作表实际范围')
            if not any(
                option['cell'] == coordinate for option in candidates.get(field, [])
            ):
                raise ValueError('字段映射必须选择该字段已识别的来源单元格')

        payload = {}
        for field, options in candidates.items():
            selected = None
            if field in mapping:
                selected = next(
                    option for option in options if option['cell'] == mapping[field]
                )
            elif len(options) == 1:
                selected = options[0]
            elif len(options) > 1:
                adaptation['notes'].append(
                    f'{LABELS[field]}有多个候选，请选择来源单元格。'
                )
                review = True
            value = selected['value'] if selected else None
            if not _empty(value):
                payload[field] = value
            adaptation['fields'].append(
                dict(
                    field=field,
                    label=LABELS[field],
                    cell=selected['cell'] if selected else None,
                    value=value,
                    candidates=options,
                )
            )
        if not payload.get('method'):
            payload['method'] = (
                'quote'
                if any(f in candidates for f in QUOTE_FIELDS)
                else 'costs' if any(f in candidates for f in COST_FIELDS) else 'ratio'
            )
        payload.setdefault('unit', '双')
        payload.setdefault('tax_basis', 'unknown')
        required = ['business_order_id']
        if payload['method'] == 'quote':
            required += ['quantity', 'unit_price', *QUOTE_FIELDS]
        elif payload['method'] == 'costs':
            required += ['quantity', 'unit_price', 'discount_per_pair', *COST_FIELDS]
        else:
            required += ['target_margin_pct']
        adaptation['missing'] = [
            field for field in required if _empty(payload.get(field))
        ]
        if payload['tax_basis'] == 'unknown':
            warnings.append('税口径尚未确认。')
        if payload['method'] == 'quote':
            warnings.append(
                '报价比例核价：生产经管计入人工，损耗计入其他生产成本；公司经管和行政单列经营费用，物流不纳入本轮预算。'
            )
        try:
            if payload['method'] == 'costs' and any(
                not _empty(payload.get(field)) for field in QUOTE_FIELDS
            ):
                message = (
                    '独立成本计算方式与报价比例字段冲突，请人工核对计算方式和单位。'
                )
                adaptation['conflicts'].append(
                    dict(
                        code='mixed_cost_basis',
                        message=message,
                        fields=[
                            field
                            for field in (*COST_FIELDS, *QUOTE_FIELDS)
                            if not _empty(payload.get(field))
                        ],
                    )
                )
                raise ValueError(message)
            if payload['method'] == 'quote' and any(
                not _empty(payload.get(field)) for field in COST_FIELDS
            ):
                message = '同时识别到报价比例和独立成本，请人工确认计算方式与成本来源；未以报价派生金额覆盖独立成本。'
                adaptation['conflicts'].append(
                    dict(
                        code='mixed_cost_basis',
                        message=message,
                        fields=[
                            field
                            for field in (*COST_FIELDS, *QUOTE_FIELDS)
                            if not _empty(payload.get(field))
                        ],
                    )
                )
                raise ValueError(message)
            baseline = normalize_baseline(payload)
        except ValueError as exc:
            baseline = payload
            warnings.append(str(exc))
            review = True
        result['baseline'] = baseline
        if not adaptation['missing'] and not review:
            result['status'] = 'ready'
        return result
    finally:
        book.close()
        cached.close()
