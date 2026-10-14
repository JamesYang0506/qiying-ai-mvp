"""Portable JSON/Excel input and traceable output bundles."""

import csv
import io
import json
import zipfile
from .training_export import export_training_bundle

HEADERS = {
    'orders': [
        'order_id',
        'product',
        'customer',
        'quantity',
        'unit',
        'unit_price',
        'discount',
        'mode',
        'scope_aligned',
        'revenue_status',
    ],
    'order_lines': [
        'order_id',
        'line_id',
        'product',
        'quantity',
        'unit',
        'unit_price',
        'discount',
    ],
    'materials': [
        'record_id',
        'order_id',
        'quantity',
        'unit_price',
        'method',
        'loss_rate',
        'loss_definition',
        'status',
        'stage',
        'period',
        'source',
    ],
    'labor': [
        'record_id',
        'order_id',
        'quantity',
        'rate',
        'status',
        'stage',
        'period',
        'source',
    ],
    'costs': [
        'record_id',
        'order_id',
        'category',
        'amount',
        'status',
        'stage',
        'period',
        'included_elsewhere',
        'source',
    ],
    'pools': [
        'pool_id',
        'category',
        'period',
        'amount',
        'total_base',
        'base_unit',
        'status',
        'direct_cost_excluded',
        'stage',
    ],
    'bases': ['pool_id', 'order_id', 'base'],
    'coverage': ['order_id', 'category', 'status'],
}


def empty_template():
    return dict(
        metadata={
            'is_demo': False,
            'currency': 'CNY',
            'cutoff_date': '2026-10-13',
            'management_enabled': False,
        },
        **{key: [] for key in HEADERS},
    )


def workbook_bytes(data):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise ValueError('Excel 功能需要 openpyxl；可先使用 JSON 模板') from None
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet('metadata')
    ws.append(['key', 'value'])
    for k, v in data['metadata'].items():
        ws.append([k, v])
    for table, headers in HEADERS.items():
        ws = wb.create_sheet(table)
        ws.append(headers)
        for row in data.get(table, []):
            ws.append(
                [
                    row.get(key, 'confirmed' if key == 'revenue_status' else None)
                    for key in headers
                ]
            )
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    for ws in wb:
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='163D35')
            cell.alignment = Alignment(vertical='center')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 24
            for cell in col[1:]:
                if isinstance(cell.value, str):
                    # Business IDs and strings are never interpreted as Excel formulas.
                    cell.data_type = 's'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def read_workbook(content):
    try:
        from openpyxl import load_workbook

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Excel 解压大小超过 50MB')
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            if all(
                s in wb.sheetnames
                for s in ('原订单用量', '补单用量', '生产周期', '成本计算')
            ):
                from .eva_source import read_eva_workbook

                return read_eva_workbook(content)
            if 'metadata' not in wb.sheetnames or 'orders' not in wb.sheetnames:
                raise ValueError('需要标准模板中的 metadata 和 orders 工作表')
            result = {'metadata': {}}
            for row in wb['metadata'].iter_rows(min_row=2, values_only=True):
                if row and row[0] is not None:
                    result['metadata'][str(row[0])] = row[1]
            for table in HEADERS:
                result[table] = []
                if table not in wb.sheetnames:
                    continue
                rows = wb[table].iter_rows(values_only=False)
                first = next(rows, None)
                if first is None:
                    continue
                headers = [
                    str(c.value).strip() if c.value is not None else '' for c in first
                ]
                if len(set(h for h in headers if h)) != len([h for h in headers if h]):
                    raise ValueError(f'{table} 存在重复列名')
                for index, cells in enumerate(rows, 2):
                    if any(c.data_type == 'f' for c in cells):
                        raise ValueError(
                            f'{table} 第 {index} 行含公式，请粘贴为数值后导入'
                        )
                    row = {h: c.value for h, c in zip(headers, cells) if h}
                    if all(v is None for v in row.values()):
                        continue
                    for key in [
                        'scope_aligned',
                        'direct_cost_excluded',
                        'included_elsewhere',
                    ]:
                        if key in row and isinstance(row[key], str):
                            v = row[key].strip().lower()
                            if v in ('true', 'false'):
                                row[key] = v == 'true'
                    row.setdefault('source', f'{table} / Excel 行 {index}')
                    result[table].append(row)
            demo = result['metadata'].get('is_demo')
            if isinstance(demo, str):
                result['metadata']['is_demo'] = demo.strip().lower() == 'true'
            return result
        finally:
            wb.close()
    except ImportError:
        raise ValueError('Excel 功能需要 openpyxl；可先使用 JSON 模板') from None
    except ValueError:
        raise
    except Exception:
        raise ValueError('无法读取 Excel，请使用未加密的标准 .xlsx 模板') from None


def export_bundle(result):
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for table in ('orders', 'costs', 'allocations', 'issues'):
            rows = []
            for record in result[table]:
                flat = {}
                for k, v in record.items():
                    if isinstance(v, dict):
                        flat.update({f'{k}.{a}': b for a, b in v.items()})
                    elif isinstance(v, list):
                        flat[k] = '; '.join(str(x) for x in v)
                    else:
                        flat[k] = v
                # Prevent spreadsheet formula execution when CSVs are opened in Excel.
                rows.append(
                    {
                        k: (
                            "'" + v
                            if isinstance(v, str)
                            and v.startswith(('=', '+', '-', '@', '\t', '\r'))
                            else v
                        )
                        for k, v in flat.items()
                    }
                )
            buf = io.StringIO(newline='')
            fields = list(dict.fromkeys(k for row in rows for k in row))
            writer = csv.DictWriter(buf, fieldnames=fields or ['empty'])
            writer.writeheader()
            writer.writerows(rows)
            z.writestr(table + '.csv', buf.getvalue().encode('utf-8-sig'))
        for name, data in [('snapshot', result['snapshot']), ('result', result)]:
            z.writestr(
                name + '.json',
                json.dumps(data, ensure_ascii=False, indent=2).encode('utf-8'),
            )
    return out.getvalue()
